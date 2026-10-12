import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment


def _write_template(path, sheet_title, columns, col_widths=None,
                    title_text=None, desc_text=None, label_row=None):
    """通用模板写入：表头样式 + 列宽。

    - 默认单表头行：表头写在第 1 行。
    - 4 行式（care/elderly）：传入 title_text / desc_text / label_row 后按 4 行结构写入：
        第 1 行：标题（跨列合并）
        第 2 行：填表说明（跨列合并）
        第 3 行：标签行，label_row 为 {列号(1-based): 标签文本}，仅填充指定列，其余留空
        第 4 行：表头（白字蓝底加粗居中）
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

    four_row = bool(title_text or desc_text or label_row)
    header_row_idx = 4 if four_row else 1

    n = len(columns)
    last_col_letter = openpyxl.utils.get_column_letter(n)

    if four_row:
        if title_text:
            c = ws.cell(row=1, column=1, value=title_text)
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center")
            ws.merge_cells(f"A1:{last_col_letter}1")
        if desc_text:
            c = ws.cell(row=2, column=1, value=desc_text)
            c.alignment = Alignment(horizontal="left", wrap_text=True)
            ws.merge_cells(f"A2:{last_col_letter}2")
        if label_row:
            for col_idx, text in label_row.items():
                ws.cell(row=3, column=col_idx, value=text)

    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=header_row_idx, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    if col_widths:
        for col_idx, width in enumerate(col_widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width
    else:
        for col_idx in range(1, n + 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 15

    wb.save(path)
