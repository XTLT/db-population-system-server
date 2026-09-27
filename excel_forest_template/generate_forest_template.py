import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from ._constants import FOREST_COLUMNS


def generate_forest_template(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "林地面积"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

    for col_idx, col_name in enumerate(FOREST_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for col_idx in range(1, len(FOREST_COLUMNS) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 18

    wb.save(path)
