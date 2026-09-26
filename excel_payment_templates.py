"""支付类数据导入模板生成模块

每种支付类型一个独立函数，由 API 统一调用。
遵循"功能模块单独写、最后统一调用"的最高级别指令。
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment


def _write_template(path, sheet_title, columns, col_widths=None):
    """通用模板写入：表头样式 + 列宽"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    if col_widths:
        for col_idx, width in enumerate(col_widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width
    else:
        for col_idx in range(1, len(columns) + 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 15

    wb.save(path)


# ── 土地补贴 ──

LAND_SUBSIDY_COLUMNS = [
    "姓名", "身份证号",
    "确权颁证面积", "补贴面积（亩）", "基础补贴标准（元/亩）",
    "其中种植油菜面积", "其中种植油菜金额", "补贴总金额（元）",
]


def generate_land_subsidy_template(path):
    _write_template(path, "地力补贴", LAND_SUBSIDY_COLUMNS)


# ── 待遇发放 ──

TREATMENT_PAYMENT_COLUMNS = [
    "所属乡镇", "所属村", "费款所属期", "支付标志", "个人编号",
    "公民身份证号码", "姓名", "银行账号", "金额", "失败原因",
]


def generate_treatment_payment_template(path):
    _write_template(path, "待遇发放", TREATMENT_PAYMENT_COLUMNS)


# ── 养老保险 ──

PENSION_INSURANCE_COLUMNS = [
    "所属乡镇", "所属村", "费款所属期", "支付标志", "个人编号",
    "享受定期待遇人员ID", "公民身份证号码", "姓名", "银行账号", "金额", "失败原因",
]


def generate_pension_insurance_template(path):
    _write_template(path, "养老保险", PENSION_INSURANCE_COLUMNS)


# ── 低保发放 ──

DIBAO_PAYMENT_COLUMNS = [
    "所属乡镇", "所属村", "费款所属期", "支付标志", "个人编号",
    "公民身份证号码", "姓名", "银行账号", "金额", "和申请人关系", "失败原因",
]


def generate_dibao_payment_template(path):
    _write_template(path, "低保发放", DIBAO_PAYMENT_COLUMNS)


# ── 特困发放 ──

TEKUN_PAYMENT_COLUMNS = [
    "所属乡镇", "所属村", "费款所属期", "支付标志", "个人编号",
    "公民身份证号码", "姓名", "银行账号", "金额", "和申请人关系", "失败原因",
]


def generate_tekun_payment_template(path):
    _write_template(path, "特困发放", TEKUN_PAYMENT_COLUMNS)
