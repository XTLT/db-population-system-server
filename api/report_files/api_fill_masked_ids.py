import os
from flask import request, jsonify

from database import get_report_file, add_report_log, get_session, Person
from utils.security import require_auth
from . import bp

# 身份证列兼容名
_ID_COLS = ('身份证号', '证件号码', '身份证', '身份证号码', 'ID')
# 姓名列兼容名
_NAME_COLS = ('姓名', '成员姓名', '户主姓名', '姓名（户主）', '姓名(户主)')


@bp.route('/<report_id>/fill_masked_ids', methods=['POST'])
@require_auth
def api_fill_masked_ids(report_id):
    """用库中数据补全报表文件中的脱敏身份证（前12位+******）。

    规则：身份证含 * 视为脱敏号，取第一个 * 前的真实部分（前 12 位）与该行姓名
    一起在库中唯一匹配（同名同前缀多人则不确定、不自动补），成功则写回完整身份证，
    其余列不变；补全后保存回原报表文件。
    """
    user = request.current_user
    report = get_report_file(report_id)
    if not report:
        return jsonify({'error': '报表不存在'}), 404

    if not os.path.exists(report['file_path']):
        return jsonify({'error': '文件不存在'}), 404

    try:
        import openpyxl
        wb = openpyxl.load_workbook(report['file_path'])
        ws = wb.active
    except Exception as e:
        return jsonify({'error': f'读取报表失败: {e}'}), 500

    id_col_idx = None
    name_col_idx = None
    for col_idx, cell in enumerate(ws[1], 1):
        val = str(cell.value or '').strip()
        if id_col_idx is None and val in _ID_COLS:
            id_col_idx = col_idx
        if name_col_idx is None and val in _NAME_COLS:
            name_col_idx = col_idx

    if id_col_idx is None:
        return jsonify({'error': '报表中未找到身份证号列（支持：身份证号/证件号码/身份证/身份证号码/ID）'}), 400
    if name_col_idx is None:
        return jsonify({'error': '报表中未找到姓名列（支持：姓名/成员姓名/户主姓名），补全需姓名+身份证前12位唯一匹配'}), 400

    session = get_session()
    filled = 0
    kept = 0
    total_mask = 0
    details = []
    try:
        for row in ws.iter_rows(min_row=2, min_col=1, max_col=ws.max_column):
            id_cell = row[id_col_idx - 1]
            id_val = str(id_cell.value or '').strip()
            if '*' not in id_val:
                continue
            total_mask += 1
            name = str(row[name_col_idx - 1].value or '').strip()
            prefix = id_val.split('*')[0].strip()
            if len(prefix) < 12:
                details.append(f"第{id_cell.row}行 {name}：脱敏号真实部分不足12位，保留")
                kept += 1
                continue
            cands = session.query(Person).filter(
                Person.name == name,
                Person.id_card.like(prefix + '%')
            ).all()
            if len(cands) == 1:
                id_cell.value = cands[0].id_card
                filled += 1
            elif len(cands) == 0:
                details.append(f"第{id_cell.row}行 {name} {id_val}：库中无此人，保留")
                kept += 1
            else:
                details.append(f"第{id_cell.row}行 {name} {id_val}：同名同前缀{len(cands)}人，无法确定，保留")
                kept += 1
        wb.save(report['file_path'])
    finally:
        session.close()

    add_report_log(
        report_id=report_id,
        user_id=user['id'],
        username=user['username'],
        operation_type='补全脱敏身份证',
        operation_detail=f'脱敏:{total_mask}, 补全:{filled}, 保留:{kept}',
        ip_address=request.remote_addr,
        client_type=request.headers.get('X-Client-Type', 'Web'),
    )

    return jsonify({'success': True, 'total_masked': total_mask, 'filled': filled, 'kept': kept, 'details': details})
