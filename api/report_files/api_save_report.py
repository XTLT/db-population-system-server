import os
from flask import request, jsonify

from database import get_report_file, add_report_log
from . import bp
from ._auth_from_query import _auth_from_query


@bp.route('/<report_id>/save', methods=['POST'])
def api_save_report(report_id):
    """保存在线编辑的报表数据"""
    user = _auth_from_query()
    if not user:
        return jsonify({'error': '未授权'}), 401

    report = get_report_file(report_id)
    if not report:
        return jsonify({'error': '报表不存在'}), 404

    if not os.path.exists(report['file_path']):
        return jsonify({'error': '文件不存在'}), 404

    data = request.get_json(silent=True) or {}
    headers = data.get('headers', [])
    rows = data.get('rows', [])

    if not headers:
        return jsonify({'error': '数据为空'}), 400

    try:
        import openpyxl
        wb = openpyxl.load_workbook(report['file_path'])
        ws = wb.active

        ws.delete_rows(1, ws.max_row)

        for col_idx, header in enumerate(headers, 1):
            ws.cell(row=1, column=col_idx, value=header)

        for row_idx, row_data in enumerate(rows, 2):
            for col_idx, val in enumerate(row_data, 1):
                ws.cell(row=row_idx, column=col_idx, value=val)

        wb.save(report['file_path'])
    except Exception as e:
        return jsonify({'error': f'保存失败: {e}'}), 500

    add_report_log(
        report_id=report_id,
        user_id=user['id'],
        username=user['username'],
        operation_type='在线编辑保存',
        operation_detail=f'行数:{len(rows)}',
        ip_address=request.remote_addr,
        client_type='Web',
    )

    return jsonify({'success': True})
