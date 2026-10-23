import os
import secrets
from flask import jsonify, send_file

from utils.security import require_auth
from utils.file_send import send_file_auto_cleanup, cleanup_temp
from config import BASE_DIR, EXCEL_TEMPLATE_COLUMNS
from . import bp


@bp.route('/template/excel')
@require_auth
def download_excel_template():
    """下载人员信息导入模板
    
    Returns:
        Excel 模板文件
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    
    tmp_path = os.path.join(BASE_DIR, f'_tmp_tpl_{secrets.token_hex(4)}.xlsx')
    
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "人口信息"
        
        # 表头样式
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        
        # 写入表头
        for col_idx, col_name in enumerate(EXCEL_TEMPLATE_COLUMNS, 1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        
        # 设置列宽
        for col_idx in range(1, len(EXCEL_TEMPLATE_COLUMNS) + 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 15
        
        wb.save(tmp_path)
        return send_file_auto_cleanup(tmp_path, as_attachment=True, 
                        download_name='人口信息导入模板.xlsx')
    except Exception as e:
        if 'tmp_path' in locals():
            cleanup_temp(tmp_path)
        return jsonify({'error': str(e)}), 400
