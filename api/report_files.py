"""报表填写管理模块 — 上传/列表/下载/在线编辑/修改日志/批量导入证件号码"""
import os
import json
from flask import Blueprint, request, jsonify, send_file, render_template_string
from werkzeug.utils import secure_filename

from database import (
    create_report_file,
    get_report_files,
    get_report_file,
    delete_report_file,
    add_report_log,
    get_report_logs,
)
from utils.security import require_auth, get_token_data
from config import REPORTS_DIR
from datetime import datetime

bp = Blueprint('report_files', __name__)

ALLOWED_EXTENSIONS = {'.xlsx', '.xls', '.et'}


def _allowed_file(filename: str) -> bool:
    ext = os.path.splitext(filename)[1].lower()
    return ext in ALLOWED_EXTENSIONS


def _auth_from_query():
    """从 query string 获取 token（用于浏览器在线编辑页面）"""
    token = request.args.get('token', '')
    if not token:
        return None
    token_data = get_token_data(token)
    if not token_data:
        return None
    from datetime import datetime as _dt
    if _dt.now() > token_data['expires_at']:
        return None
    return token_data['user']


@bp.route('/upload', methods=['POST'])
@require_auth
def api_upload_report():
    """上传 Excel 报表文件"""
    user = request.current_user

    if 'file' not in request.files:
        return jsonify({'error': '未选择文件'}), 400

    file = request.files['file']
    if not file.filename:
        return jsonify({'error': '未选择文件'}), 400

    original_name = file.filename
    if not _allowed_file(original_name):
        return jsonify({'error': '仅支持 .xlsx .xls .et 格式'}), 400

    name = request.form.get('name', '').strip()
    if not name:
        name = original_name

    file_data = file.read()
    file_size = len(file_data)

    stored_filename = f"report_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{secure_filename(original_name)}"
    file_path = os.path.join(REPORTS_DIR, stored_filename)

    with open(file_path, 'wb') as f:
        f.write(file_data)

    report = create_report_file(
        name=name,
        original_name=original_name,
        stored_filename=stored_filename,
        file_path=file_path,
        file_size=file_size,
        user_id=user['id'],
        username=user['username'],
    )

    add_report_log(
        report_id=report['report_id'],
        user_id=user['id'],
        username=user['username'],
        operation_type='上传报表',
        operation_detail=f'文件:{original_name}, 名称:{name}',
        ip_address=request.remote_addr,
        client_type='Desktop',
    )

    return jsonify({'success': True, 'report': report})


@bp.route('', methods=['GET'])
@require_auth
def api_list_reports():
    """获取报表文件列表"""
    reports = get_report_files()
    return jsonify(reports)


@bp.route('/<report_id>/download')
@require_auth
def api_download_report(report_id):
    """下载报表文件"""
    report = get_report_file(report_id)
    if not report:
        return jsonify({'error': '报表不存在'}), 404

    if not os.path.exists(report['file_path']):
        return jsonify({'error': '文件不存在'}), 404

    user = request.current_user
    add_report_log(
        report_id=report_id,
        user_id=user['id'],
        username=user['username'],
        operation_type='下载报表',
        operation_detail=f'文件:{report["original_name"]}',
        ip_address=request.remote_addr,
        client_type='Desktop',
    )

    return send_file(
        report['file_path'],
        as_attachment=True,
        download_name=report['original_name'],
    )


@bp.route('/<report_id>/logs')
@require_auth
def api_get_report_logs(report_id):
    """获取报表操作日志"""
    report = get_report_file(report_id)
    if not report:
        return jsonify({'error': '报表不存在'}), 404

    logs = get_report_logs(report_id)
    return jsonify({'db_logs': logs, 'file_logs': []})


@bp.route('/<report_id>/import_ids', methods=['POST'])
@require_auth
def api_import_ids(report_id):
    """批量导入证件号码到报表"""
    user = request.current_user
    report = get_report_file(report_id)
    if not report:
        return jsonify({'error': '报表不存在'}), 404

    if not os.path.exists(report['file_path']):
        return jsonify({'error': '文件不存在'}), 404

    data = request.get_json(silent=True) or {}
    ids = data.get('ids', [])
    if not ids:
        return jsonify({'error': '没有可导入的数据'}), 400

    id_map = {}
    for item in ids:
        id_card = str(item.get('id_card', '')).strip()
        if id_card:
            id_map[id_card] = item

    try:
        import openpyxl
        wb = openpyxl.load_workbook(report['file_path'])
        ws = wb.active
    except Exception as e:
        return jsonify({'error': f'读取报表失败: {e}'}), 500

    id_col_idx = None
    phone_col_idx = None
    social_cols = {}

    for col_idx, cell in enumerate(ws[1], 1):
        val = str(cell.value or '').strip()
        if val in ('身份证号', '证件号码', '身份证', '身份证号码', 'ID'):
            id_col_idx = col_idx
        elif val in ('电话', '手机号', '手机', '联系电话', '联系方式'):
            phone_col_idx = col_idx
        elif val in ('社保卡1', '社保卡号1', '社会保障卡1'):
            social_cols[1] = col_idx
        elif val in ('社保卡2', '社保卡号2', '社会保障卡2'):
            social_cols[2] = col_idx
        elif val in ('社保卡3', '社保卡号3', '社会保障卡3'):
            social_cols[3] = col_idx

    if id_col_idx is None:
        return jsonify({'error': '报表中未找到身份证号列'}), 400

    updated = 0
    not_found = 0

    for row in ws.iter_rows(min_row=2, min_col=1, max_col=ws.max_column):
        id_cell = row[id_col_idx - 1]
        id_val = str(id_cell.value or '').strip()
        if not id_val:
            continue

        if id_val in id_map:
            item = id_map[id_val]
            if phone_col_idx and item.get('phone'):
                ws.cell(row=id_cell.row, column=phone_col_idx, value=item['phone'])
            for key, col_idx in social_cols.items():
                field = f'social{key}'
                if item.get(field):
                    ws.cell(row=id_cell.row, column=col_idx, value=item[field])
            updated += 1
        else:
            not_found += 1

    wb.save(report['file_path'])

    add_report_log(
        report_id=report_id,
        user_id=user['id'],
        username=user['username'],
        operation_type='批量导入证件号码',
        operation_detail=f'更新:{updated}, 未匹配:{not_found}',
        ip_address=request.remote_addr,
        client_type='Desktop',
    )

    return jsonify({'success': True, 'updated': updated, 'not_found': not_found})


@bp.route('/<report_id>/edit')
def api_edit_report(report_id):
    """在线编辑报表页面（浏览器打开）"""
    user = _auth_from_query()
    if not user:
        return '<html><body><h2>登录已过期，请重新登录后重试</h2></body></html>', 401

    report = get_report_file(report_id)
    if not report:
        return '<html><body><h2>报表不存在</h2></body></html>', 404

    if not os.path.exists(report['file_path']):
        return '<html><body><h2>报表文件不存在</h2></body></html>', 404

    token = request.args.get('token', '')
    return _render_editor_html(report_id, report['name'], token)


@bp.route('/<report_id>/data')
def api_get_report_data(report_id):
    """获取报表数据（JSON 格式，供在线编辑器使用）"""
    user = _auth_from_query()
    if not user:
        return jsonify({'error': '未授权'}), 401

    report = get_report_file(report_id)
    if not report:
        return jsonify({'error': '报表不存在'}), 404

    if not os.path.exists(report['file_path']):
        return jsonify({'error': '文件不存在'}), 404

    try:
        import openpyxl
        wb = openpyxl.load_workbook(report['file_path'], data_only=True)
        ws = wb.active
    except Exception as e:
        return jsonify({'error': f'读取失败: {e}'}), 500

    headers = []
    rows = []
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
        if row_idx == 1:
            headers = [str(cell or '') for cell in row]
        else:
            if all(cell is None for cell in row):
                continue
            rows.append([str(cell) if cell is not None else '' for cell in row])

    return jsonify({
        'sheet_name': ws.title,
        'headers': headers,
        'rows': rows,
    })


@bp.route('/<report_id>/save', methods=['POST'])
def api_save_report(report_id):
    """保存在线编辑的报表数据"""
    user = _auth_from_query()
    if not user:
        return jsonify({'error': '未授权'}), 401

    report = get_report_file(report_id)
    if not report:
        return jsonify({'error': '报表不存在'}), 404

    if not os.path.exists(report['file_path']):
        return jsonify({'error': '文件不存在'}), 404

    data = request.get_json(silent=True) or {}
    headers = data.get('headers', [])
    rows = data.get('rows', [])

    if not headers:
        return jsonify({'error': '数据为空'}), 400

    try:
        import openpyxl
        wb = openpyxl.load_workbook(report['file_path'])
        ws = wb.active

        ws.delete_rows(1, ws.max_row)

        for col_idx, header in enumerate(headers, 1):
            ws.cell(row=1, column=col_idx, value=header)

        for row_idx, row_data in enumerate(rows, 2):
            for col_idx, val in enumerate(row_data, 1):
                ws.cell(row=row_idx, column=col_idx, value=val)

        wb.save(report['file_path'])
    except Exception as e:
        return jsonify({'error': f'保存失败: {e}'}), 500

    add_report_log(
        report_id=report_id,
        user_id=user['id'],
        username=user['username'],
        operation_type='在线编辑保存',
        operation_detail=f'行数:{len(rows)}',
        ip_address=request.remote_addr,
        client_type='Web',
    )

    return jsonify({'success': True})


@bp.route('/<report_id>', methods=['DELETE'])
@require_auth
def api_delete_report(report_id):
    """删除报表"""
    user = request.current_user
    report = get_report_file(report_id)
    if not report:
        return jsonify({'error': '报表不存在'}), 404

    if os.path.exists(report['file_path']):
        try:
            os.remove(report['file_path'])
        except OSError:
            pass

    delete_report_file(report_id)

    add_report_log(
        report_id=report_id,
        user_id=user['id'],
        username=user['username'],
        operation_type='删除报表',
        operation_detail=f'文件:{report["original_name"]}',
        ip_address=request.remote_addr,
        client_type='Desktop',
    )

    return jsonify({'success': True})


def _render_editor_html(report_id: str, report_name: str, token: str) -> str:
    """渲染在线编辑器 HTML 页面"""
    return render_template_string(EDITOR_HTML_TEMPLATE,
                                  report_id=report_id,
                                  report_name=report_name,
                                  token=token)


EDITOR_HTML_TEMPLATE = '''
<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>在线编辑 - {{ report_name }}</title>
<style>
* { margin: 0; padding: 0; box-sizing: border-box; }
body { font-family: "Microsoft YaHei", sans-serif; background: #f5f5f5; }
.header { background: #4472C4; color: white; padding: 12px 20px; display: flex; justify-content: space-between; align-items: center; }
.header h1 { font-size: 18px; }
.toolbar { padding: 10px 20px; background: white; border-bottom: 1px solid #ddd; display: flex; gap: 10px; }
.btn { padding: 6px 16px; border: none; border-radius: 4px; cursor: pointer; font-size: 14px; }
.btn-save { background: #4472C4; color: white; }
.btn-save:hover { background: #3656b8; }
.btn-cancel { background: #e0e0e0; }
.btn-cancel:hover { background: #d0d0d0; }
.table-container { overflow: auto; max-height: calc(100vh - 120px); padding: 10px; }
table { border-collapse: collapse; background: white; }
th, td { border: 1px solid #ccc; padding: 4px 8px; min-width: 80px; max-width: 300px; }
th { background: #4472C4; color: white; position: sticky; top: 0; font-weight: bold; }
td { white-space: pre-wrap; word-break: break-all; }
td[contenteditable="true"]:focus { background: #e6f2ff; outline: 2px solid #4472C4; }
.status { padding: 8px 20px; font-size: 13px; color: #666; }
.status .success { color: #4CAF50; }
.status .error { color: #f44336; }
.loading { text-align: center; padding: 40px; font-size: 16px; color: #999; }
</style>
</head>
<body>
<div class="header">
  <h1>📋 {{ report_name }}</h1>
</div>
<div class="toolbar">
  <button class="btn btn-save" onclick="saveData()">💾 保存</button>
  <button class="btn btn-cancel" onclick="window.close()">关闭</button>
  <span id="status" class="status"></span>
</div>
<div id="content" class="table-container">
  <div class="loading">正在加载报表数据...</div>
</div>
<script>
const REPORT_ID = '{{ report_id }}';
const TOKEN = '{{ token }}';

function showStatus(msg, type) {
  const el = document.getElementById('status');
  el.innerHTML = '<span class="' + (type||'') + '">' + msg + '</span>';
}

async function loadData() {
  try {
    const resp = await fetch('/api/reports/' + REPORT_ID + '/data?token=' + TOKEN);
    if (!resp.ok) {
      const err = await resp.json().catch(() => ({error: 'HTTP ' + resp.status}));
      throw new Error(err.error || '加载失败');
    }
    const data = await resp.json();
    renderTable(data.headers, data.rows);
    showStatus('已加载 ' + data.rows.length + ' 行数据', 'success');
  } catch(e) {
    document.getElementById('content').innerHTML = '<div class="loading" style="color:#f44336;">加载失败: ' + e.message + '</div>';
  }
}

function renderTable(headers, rows) {
  let html = '<table><thead><tr>';
  headers.forEach(h => { html += '<th>' + escapeHtml(h) + '</th>'; });
  html += '</tr></thead><tbody>';
  rows.forEach((row, ri) => {
    html += '<tr>';
    headers.forEach((h, ci) => {
      const val = row[ci] !== undefined ? escapeHtml(row[ci]) : '';
      html += '<td contenteditable="true" data-row="' + ri + '" data-col="' + ci + '">' + val + '</td>';
    });
    html += '</tr>';
  });
  html += '</tbody></table>';
  document.getElementById('content').innerHTML = html;
}

function escapeHtml(s) {
  if (s === null || s === undefined) return '';
  return String(s).replace(/&/g,'&amp;').replace(/</g,'&lt;').replace(/>/g,'&gt;').replace(/"/g,'&quot;');
}

async function saveData() {
  const table = document.querySelector('table');
  if (!table) return;

  const headers = [];
  table.querySelectorAll('thead th').forEach(th => {
    headers.push(th.textContent);
  });

  const rows = [];
  table.querySelectorAll('tbody tr').forEach(tr => {
    const row = [];
    tr.querySelectorAll('td').forEach(td => {
      row.push(td.textContent);
    });
    rows.push(row);
  });

  showStatus('正在保存...', '');

  try {
    const resp = await fetch('/api/reports/' + REPORT_ID + '/save?token=' + TOKEN, {
      method: 'POST',
      headers: {'Content-Type': 'application/json'},
      body: JSON.stringify({headers: headers, rows: rows})
    });
    if (!resp.ok) {
      const err = await resp.json().catch(() => ({error: 'HTTP ' + resp.status}));
      throw new Error(err.error || '保存失败');
    }
    showStatus('保存成功 (' + rows.length + ' 行)', 'success');
  } catch(e) {
    showStatus('保存失败: ' + e.message, 'error');
  }
}

loadData();
</script>
</body>
</html>
'''
