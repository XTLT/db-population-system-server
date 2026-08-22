import sys
import os
import json
import secrets
import functools
import shutil
from datetime import datetime, timedelta

from flask import Flask, jsonify, request, send_from_directory, send_file
from flask_cors import CORS
from sqlalchemy import func
from database import init_db, get_session, Person, StudentEducation, search_by_name, search_persons_advanced, get_family_members, get_person_count, update_person, add_person, delete_person, migrate_person, get_student_education, get_all_students, search_students, create_user, get_all_users, delete_user, User, OperationLog, add_operation_log, get_operation_logs, get_operation_types, get_operation_log_stats, validate_password, ReportOperationLog, add_report_log, get_report_logs, get_all_report_logs, create_notification, get_notifications, mark_notification_read, get_unread_count, get_person_trend_data, get_person_distribution, get_person_comparison, create_custom_report, get_custom_reports, get_custom_report, update_custom_report, delete_custom_report, generate_report_data
from utils import calculate_age, find_existing_photo, get_photo_base_name, clean_filename
from config import GRADE_DISPLAY, TOKEN_EXPIRY_HOURS, PHOTO_EXTENSIONS, BASE_DIR, LOGIN_LOCK_MINUTES
import config
from update_checker import check_update, download_update, verify_file, update_local_version, apply_update
from ssl_utils import get_cert_paths

IMAGE_MAGIC_BYTES = {
    b'\xff\xd8\xff': 'jpeg',
    b'\x89PNG\r\n\x1a\n': 'png',
    b'GIF87a': 'gif',
    b'GIF89a': 'gif',
    b'RIFF': 'webp',
    b'BM': 'bmp',
}

def validate_image_content(file_obj):
    try:
        file_obj.seek(0)
        header = file_obj.read(16)
        file_obj.seek(0)
        for magic, img_type in IMAGE_MAGIC_BYTES.items():
            if header.startswith(magic):
                return img_type
        return None
    except Exception:
        return None

if getattr(sys, 'frozen', False):
    _MEIPASS = sys._MEIPASS
    _search_dirs = [
        os.path.join(_MEIPASS, 'server', 'web_static'),
        os.path.join(_MEIPASS, 'web_static'),
    ]
else:
    _search_dirs = [os.path.join(BASE_DIR, 'web_static')]

STATIC_DIR = None
for _d in _search_dirs:
    if os.path.isdir(_d):
        STATIC_DIR = _d
        break
if STATIC_DIR is None:
    STATIC_DIR = _search_dirs[0]

app = Flask(__name__)

app.secret_key = os.environ.get('SECRET_KEY') or secrets.token_hex(32)


_cors_origins = os.environ.get('CORS_ORIGINS', 'http://localhost:5000,http://127.0.0.1:5000')
CORS_ORIGINS = [origin.strip() for origin in _cors_origins.split(',') if origin.strip()]

CORS(app, resources={
    r"/api/*": {
        "origins": CORS_ORIGINS if CORS_ORIGINS else "*",
        "methods": ["GET", "POST", "PUT", "DELETE", "OPTIONS"],
        "allow_headers": ["Content-Type", "Authorization"],
        "supports_credentials": True,
    }
})

_tokens = {}

_login_attempts = {}
LOGIN_RATE_LIMIT = 10
LOGIN_RATE_WINDOW = 60

_sensitive_op_attempts = {}
SENSITIVE_OP_RATE_LIMIT = 5
SENSITIVE_OP_RATE_WINDOW = 60


def check_sensitive_op_rate_limit(ip: str, op_type: str = 'default') -> tuple:
    """检查敏感操作速率限制
    
    用于密码修改、重置等敏感操作，防止暴力攻击
    
    Args:
        ip: 客户端IP
        op_type: 操作类型（可选，用于区分不同操作的限额）
        
    Returns:
        (is_allowed, remaining_seconds)
    """
    import time
    now = time.time()
    key = f"{ip}:{op_type}"
    
    # 清理过期记录
    if key in _sensitive_op_attempts:
        _sensitive_op_attempts[key] = [
            ts for ts in _sensitive_op_attempts[key]
            if now - ts < SENSITIVE_OP_RATE_WINDOW
        ]
    
    if key not in _sensitive_op_attempts:
        _sensitive_op_attempts[key] = []
    
    if len(_sensitive_op_attempts[key]) >= SENSITIVE_OP_RATE_LIMIT:
        oldest = _sensitive_op_attempts[key][0]
        remaining = int(SENSITIVE_OP_RATE_WINDOW - (now - oldest))
        return False, max(remaining, 1)
    
    return True, 0


def record_sensitive_op_attempt(ip: str, op_type: str = 'default'):
    """记录敏感操作尝试"""
    import time
    key = f"{ip}:{op_type}"
    if key not in _sensitive_op_attempts:
        _sensitive_op_attempts[key] = []
    _sensitive_op_attempts[key].append(time.time())


def check_login_rate_limit(ip: str) -> tuple:
    """检查登录速率限制
    
    Returns:
        (is_allowed, remaining_seconds)
    """
    import time
    now = time.time()
    
    # 清理过期记录
    if ip in _login_attempts:
        _login_attempts[ip] = [
            (ts, success) for ts, success in _login_attempts[ip]
            if now - ts < LOGIN_RATE_WINDOW
        ]
    
    if ip not in _login_attempts:
        _login_attempts[ip] = []
    
    # 计算失败次数
    failed_count = sum(1 for ts, success in _login_attempts[ip] if not success)
    
    if failed_count >= LOGIN_RATE_LIMIT:
        oldest = min(ts for ts, success in _login_attempts[ip] if not success)
        remaining = int(LOGIN_RATE_WINDOW - (now - oldest))
        return False, max(remaining, 1)
    
    return True, 0


def record_login_attempt(ip: str, success: bool):
    """记录登录尝试"""
    import time
    if ip not in _login_attempts:
        _login_attempts[ip] = []
    _login_attempts[ip].append((time.time(), success))


def hash_password(password: str) -> str:
    import hashlib
    salt = secrets.token_hex(16)
    password_sha256 = hashlib.sha256(password.encode()).hexdigest()
    password_hash = hashlib.sha256((salt + password_sha256).encode()).hexdigest()
    return f"{salt}:{password_hash}"


def verify_password(password: str, password_hash: str) -> bool:
    import hashlib
    try:
        salt, stored_hash = password_hash.split(":")
        
        password_sha256 = hashlib.sha256(password.encode()).hexdigest()
        new_format_hash = hashlib.sha256((salt + password_sha256).encode()).hexdigest()
        if new_format_hash == stored_hash:
            return True
        
        old_format_hash = hashlib.sha256((salt + password).encode()).hexdigest()
        if old_format_hash == stored_hash:
            return True
        
        return False
    except Exception:
        return False


def _get_json():
    data = request.json
    if data is None:
        return None, (jsonify({'error': '请求数据格式错误'}), 400)
    return data, None


def safe_error(message: str, status: int = 400, log_details: bool = True):
    if log_details:
        import traceback
        print(f"[错误] {message}")
        if os.environ.get('FLASK_DEBUG') == '1':
            traceback.print_exc()
    
    return jsonify({'error': message}), status


_backup_download_times = {}

_id_card_query_times = {}
ID_CARD_QUERY_LIMIT = 30
ID_CARD_QUERY_WINDOW = 60


def check_id_card_rate_limit(person_id: str) -> tuple:
    ip = request.remote_addr
    now = datetime.now()
    
    if ip in _id_card_query_times:
        _id_card_query_times[ip] = [
            (ts, pid) for ts, pid in _id_card_query_times[ip]
            if (now - ts).total_seconds() < ID_CARD_QUERY_WINDOW
        ]
    
    if ip not in _id_card_query_times:
        _id_card_query_times[ip] = []
    
    query_count = len(_id_card_query_times[ip])
    
    if query_count >= ID_CARD_QUERY_LIMIT:
        oldest = _id_card_query_times[ip][0][0]
        remaining = ID_CARD_QUERY_WINDOW - (now - oldest).total_seconds()
        return False, int(remaining)
    
    _id_card_query_times[ip].append((now, person_id))
    return True, 0


def safe_id_card_query(session, person_id: str):
    from database import Person
    
    p = session.query(Person).filter(Person.id == person_id).first()
    if p:
        return p, None
    
    if len(person_id) == 18 and person_id.isdigit():
        is_allowed, remaining = check_id_card_rate_limit(person_id)
        if not is_allowed:
            return None, (jsonify({
                'error': f'查询过于频繁，请等待{remaining}秒后重试'
            }), 429)
        
        p = session.query(Person).filter(Person.id_card == person_id).first()
        if p:
            return p, None
    
    return None, (jsonify({'error': 'Not found'}), 404)


@app.after_request
def hide_server_info(response):
    response.headers.pop('Server', None)
    
    response.headers.pop('ETag', None)
    
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['X-Frame-Options'] = 'DENY'
    response.headers['X-XSS-Protection'] = '1; mode=block'
    response.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate'
    response.headers['Pragma'] = 'no-cache'
    
    response.headers['Referrer-Policy'] = 'strict-origin-when-cross-origin'
    
    response.headers['Content-Security-Policy'] = "default-src 'self'; script-src 'self' 'unsafe-inline'; style-src 'self' 'unsafe-inline'; img-src 'self' data:; font-src 'self'; connect-src 'self'; frame-ancestors 'none';"
    
    response.headers['Permissions-Policy'] = 'geolocation=(), microphone=(), camera=(), payment=(), usb=(), magnetometer=(), gyroscope=(), accelerometer=()'
    
    cert_file, key_file = get_cert_paths()
    if cert_file and os.path.exists(cert_file):
        response.headers['Strict-Transport-Security'] = 'max-age=31536000; includeSubDomains'
    
    if os.environ.get('FLASK_HTTPS_REDIRECT', '1') == '1':
        if not request.is_secure and request.headers.get('X-Forwarded-Proto', 'http') != 'https':
            if cert_file and os.path.exists(cert_file):
                url = request.url.replace('http://', 'https://', 1)
                return response.redirect(url, code=301)
    return response


def require_auth(f):
    @functools.wraps(f)
    def decorated(*args, **kwargs):
        token = request.headers.get('Authorization', '').replace('Bearer ', '')
        if not token or token not in _tokens:
            return jsonify({'error': '未授权访问'}), 401
        token_data = _tokens[token]
        if datetime.now() > token_data['expires_at']:
            del _tokens[token]
            return jsonify({'error': '登录已过期，请重新登录'}), 401
        request.current_user = token_data['user']
        return f(*args, **kwargs)
    return decorated


@app.route('/')
def index():
    return send_from_directory(STATIC_DIR, 'index.html')


@app.route('/style.css')
def style():
    return send_from_directory(STATIC_DIR, 'style.css')


@app.route('/app.js')
def script():
    return send_from_directory(STATIC_DIR, 'app.js')


DEFAULT_ADMIN_USERNAME = "admin"
DEFAULT_ADMIN_PASSWORD = "admin123"


def _ensure_default_admin():
    session = get_session()
    try:
        existing = session.query(User).filter(User.username == DEFAULT_ADMIN_USERNAME).first()
        if not existing:
            import hashlib
            salt = secrets.token_hex(16)
            password_sha256 = hashlib.sha256(DEFAULT_ADMIN_PASSWORD.encode()).hexdigest()
            password_hash = hashlib.sha256((salt + password_sha256).encode()).hexdigest()
            user = User(
                id=f"user_default_{secrets.token_hex(8)}",
                username=DEFAULT_ADMIN_USERNAME,
                password_hash=f"{salt}:{password_hash}",
                real_name="初始管理员",
                role="admin",
                is_active="是",
                is_default="是",
                created_at=datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            )
            session.add(user)
            session.commit()
    except Exception as e:
        session.rollback()
    finally:
        session.close()


@app.route('/api/status')
def api_status():
    _ensure_default_admin()
    users = get_all_users()
    has_default = any(u.is_default == "是" for u in users)
    has_real_admin = any(u.is_default != "是" and u.role == "admin" for u in users)
    return jsonify({
        'initialized': len(users) > 0,
        'has_default_admin': has_default,
        'need_create_admin': has_default and not has_real_admin
    })


@app.route('/api/health')
def api_health():
    return jsonify({'status': 'ok', 'healthy': True})


@app.route('/api/init', methods=['POST'])
def api_init():
    data, err = _get_json()
    if err:
        return err
    username = data.get('username', '')
    password = data.get('password', '')
    real_name = data.get('real_name', '')
    users = get_all_users()
    if users:
        return jsonify({'error': '系统已初始化'}), 400
    valid, msg = validate_password(password)
    if not valid:
        return jsonify({'error': msg}), 400
    try:
        create_user(username, password, real_name, "admin")
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 400


@app.route('/api/login', methods=['POST'])
def api_login():
    data, err = _get_json()
    if err:
        return err
    
    ip_address = request.remote_addr
    
    is_allowed, remaining = check_login_rate_limit(ip_address)
    if not is_allowed:
        return jsonify({
            'success': False, 
            'error': f'登录尝试过于频繁，请{remaining}秒后重试'
        }), 429
    
    username = data.get('username', '')
    password = data.get('password', '')
    password_hashed = data.get('password_hashed', False)

    session = get_session()
    try:
        user_obj = session.query(User).filter(User.username == username, User.is_active == "是").first()
        if not user_obj:
            record_login_attempt(ip_address, False)
            return jsonify({'success': False, 'error': '用户名或密码错误'}), 401

        if user_obj.locked_until:
            try:
                lock_time = datetime.strptime(user_obj.locked_until, "%Y-%m-%d %H:%M:%S")
                if datetime.now() < lock_time:
                    remaining = int((lock_time - datetime.now()).total_seconds() / 60) + 1
                    return jsonify({'success': False, 'error': f'密码错误次数过多，账户已锁定，请{remaining}分钟后重试'}), 429
                user_obj.locked_until = ""
                user_obj.login_attempts = "0"
            except Exception:
                user_obj.locked_until = ""
                user_obj.login_attempts = "0"

        import hashlib
        
        try:
            salt, stored_hash = user_obj.password_hash.split(":")
        except (ValueError, AttributeError):
            return jsonify({'success': False, 'error': '用户数据格式错误'}), 500
        
        verified = False
        need_migrate = False
        new_hash = None
        
        if password_hashed:
            new_hash = hashlib.sha256((salt + password).encode()).hexdigest()
            if new_hash == stored_hash:
                verified = True
        else:
            old_hash = hashlib.sha256((salt + password).encode()).hexdigest()
            if old_hash == stored_hash:
                verified = True
                need_migrate = True
                password_sha256 = hashlib.sha256(password.encode()).hexdigest()
                new_hash = hashlib.sha256((salt + password_sha256).encode()).hexdigest()
            else:
                password_sha256 = hashlib.sha256(password.encode()).hexdigest()
                new_hash = hashlib.sha256((salt + password_sha256).encode()).hexdigest()
                if new_hash == stored_hash:
                    verified = True
        
        if not verified:
            attempts = int(user_obj.login_attempts or "0") + 1
            user_obj.login_attempts = str(attempts)
            if attempts >= 5:
                user_obj.locked_until = (datetime.now() + timedelta(minutes=LOGIN_LOCK_MINUTES)).strftime("%Y-%m-%d %H:%M:%S")
            session.commit()
            record_login_attempt(ip_address, False)  # 记录失败尝试
            return jsonify({'success': False, 'error': '用户名或密码错误'}), 401

        if need_migrate and new_hash:
            user_obj.password_hash = f"{salt}:{new_hash}"
            session.commit()

        is_default = user_obj.is_default == "是"

        user_obj.login_attempts = "0"
        user_obj.locked_until = ""
        user_obj.last_login = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        session.commit()
        
        record_login_attempt(ip_address, True)  # 记录成功登录

        token = secrets.token_hex(32)
        _tokens[token] = {
            'user': {
                'id': user_obj.id,
                'username': user_obj.username,
                'real_name': user_obj.real_name,
                'role': user_obj.role,
                'is_default': user_obj.is_default
            },
            'expires_at': datetime.now() + timedelta(hours=TOKEN_EXPIRY_HOURS)
        }
        add_operation_log(user_obj.id, user_obj.username, "登录", "", ip_address, "Web")

        return jsonify({
            'success': True,
            'token': token,
            'user': {
                'username': user_obj.username,
                'real_name': user_obj.real_name,
                'role': user_obj.role
            },
            'is_default_user': is_default,
            'require_password_change': is_default
        })
    except Exception as e:
        session.rollback()
        import traceback
        traceback.print_exc()
        return jsonify({'success': False, 'error': f'服务器错误: {str(e)}'}), 500
    finally:
        session.close()


@app.route('/api/logout', methods=['POST'])
@require_auth
def api_logout():
    token = request.headers.get('Authorization', '').replace('Bearer ', '')
    if token in _tokens:
        del _tokens[token]
    return jsonify({'success': True})


@app.route('/api/create-admin', methods=['POST'])
@require_auth
def api_create_admin():
    data, err = _get_json()
    if err:
        return err
    username = data.get('username', '')
    password = data.get('password', '')
    real_name = data.get('real_name', '')

    session = get_session()
    try:
        has_real_admin = session.query(User).filter(
            User.is_default != "是",
            User.role == "admin"
        ).first()
        if has_real_admin:
            return jsonify({'error': '系统已初始化，请使用管理员账户创建新用户'}), 400

        if request.current_user.get('username') != DEFAULT_ADMIN_USERNAME:
            return jsonify({'error': '权限不足，仅允许默认管理员执行此操作'}), 403

        valid, msg = validate_password(password)
        if not valid:
            return jsonify({'error': msg}), 400

        existing = session.query(User).filter(User.username == username).first()
        if existing:
            return jsonify({'error': '用户名已存在'}), 400

        import hashlib
        salt = secrets.token_hex(16)
        password_sha256 = hashlib.sha256(password.encode()).hexdigest()
        password_hash = hashlib.sha256((salt + password_sha256).encode()).hexdigest()
        user = User(
            id=f"user_{secrets.token_hex(8)}",
            username=username,
            password_hash=f"{salt}:{password_hash}",
            real_name=real_name,
            role="admin",
            is_default="否",
            created_at=datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        )
        session.add(user)

        default_user = session.query(User).filter(User.is_default == "是").first()
        if default_user:
            default_user.is_active = "否"

        session.commit()

        add_operation_log(user.id, user.username, "创建管理员账户(首次初始化)", "", "", "Web")

        return jsonify({'success': True})
    except Exception as e:
        session.rollback()
        return jsonify({'error': str(e)}), 400
    finally:
        session.close()


@app.route('/api/users')
@require_auth
def api_users():
    if request.current_user.get('role') != 'admin':
        return jsonify({'error': '权限不足'}), 403
    users = get_all_users()
    return jsonify({
        'users': [{
            'id': u.id,
            'username': u.username,
            'real_name': u.real_name,
            'role': u.role,
            'is_active': u.is_active,
            'is_default': u.is_default,
            'created_at': u.created_at or '',
            'last_login': u.last_login or '',
            'login_attempts': u.login_attempts or 0
        } for u in users]
    })


@app.route('/api/users', methods=['POST'])
@require_auth
def api_add_user():
    if request.current_user.get('role') != 'admin':
        return jsonify({'error': '权限不足'}), 403
    data, err = _get_json()
    if err:
        return err
    try:
        create_user(
            data['username'],
            data['password'],
            data.get('real_name', ''),
            data.get('role', 'user')
        )
        add_operation_log(request.current_user['id'], request.current_user['username'], "添加用户", data['username'], request.remote_addr, "Web")
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 400


@app.route('/api/users/<user_id>', methods=['DELETE'])
@require_auth
def api_delete_user(user_id):
    if request.current_user.get('role') != 'admin':
        return jsonify({'error': '权限不足'}), 403
    delete_user(user_id)
    add_operation_log(request.current_user['id'], request.current_user['username'], "删除用户", user_id, request.remote_addr, "Web")
    return jsonify({'success': True})


@app.route('/api/users/<user_id>', methods=['PUT'])
@require_auth
def api_update_user(user_id):
    is_admin = request.current_user.get('role') == 'admin'
    is_self = request.current_user.get('id') == user_id
    if not is_admin and not is_self:
        return jsonify({'error': '权限不足'}), 403
    data, err = _get_json()
    if err:
        return err
    session = get_session()
    try:
        user = session.query(User).filter(User.id == user_id).first()
        if not user:
            return jsonify({'error': '用户不存在'}), 404
        if 'real_name' in data:
            user.real_name = data['real_name']
        if is_admin and 'role' in data:
            user.role = data['role']
        if is_admin and 'is_active' in data:
            user.is_active = data['is_active']
        if 'username' in data and data['username'] != user.username:
            existing = session.query(User).filter(User.username == data['username']).first()
            if existing:
                return jsonify({'error': '用户名已存在'}), 400
            user.username = data['username']
        if 'password' in data and data['password']:
            import hashlib
            if not is_admin and 'old_password' in data:
                salt, stored_hash = user.password_hash.split(":")
                old_sha256 = hashlib.sha256(data['old_password'].encode()).hexdigest()
                new_verify = hashlib.sha256((salt + old_sha256).encode()).hexdigest()
                old_verify = hashlib.sha256((salt + data['old_password']).encode()).hexdigest()
                if new_verify != stored_hash and old_verify != stored_hash:
                    return jsonify({'error': '原密码错误'}), 400
            salt = secrets.token_hex(16)
            password_sha256 = hashlib.sha256(data['password'].encode()).hexdigest()
            password_hash = hashlib.sha256((salt + password_sha256).encode()).hexdigest()
            user.password_hash = f"{salt}:{password_hash}"
        session.commit()
        return jsonify({'success': True})
    except Exception as e:
        session.rollback()
        return jsonify({'error': str(e)}), 400
    finally:
        session.close()


@app.route('/api/users/<user_id>/reset-password', methods=['POST'])
@require_auth
def api_reset_password(user_id):
    if request.current_user.get('role') != 'admin':
        return jsonify({'error': '权限不足'}), 403
    
    # 检查速率限制
    ip = request.remote_addr
    is_allowed, remaining = check_sensitive_op_rate_limit(ip, 'reset_password')
    if not is_allowed:
        return jsonify({'error': f'操作过于频繁，请{remaining}秒后重试'}), 429
    
    data, err = _get_json()
    if err:
        return err
    
    new_password = data.get('password', '')
    
    if not new_password or len(new_password) < 6:
        return jsonify({'error': '密码长度至少6位'}), 400
    
    # 记录操作尝试
    record_sensitive_op_attempt(ip, 'reset_password')
    
    session = get_session()
    try:
        user = session.query(User).filter(User.id == user_id).first()
        if not user:
            return jsonify({'error': '用户不存在'}), 404
        
        import hashlib
        salt = secrets.token_hex(16)
        password_sha256 = hashlib.sha256(new_password.encode()).hexdigest()
        password_hash = hashlib.sha256((salt + password_sha256).encode()).hexdigest()
        user.password_hash = f"{salt}:{password_hash}"
        user.login_attempts = "0"
        user.locked_until = ""
        
        session.commit()
        
        add_operation_log(
            request.current_user['id'],
            request.current_user['username'],
            "重置密码",
            user.username,
            request.remote_addr,
            "Web"
        )
        
        return jsonify({'success': True, 'message': '密码重置成功'})
    except Exception as e:
        session.rollback()
        return jsonify({'error': str(e)}), 500
    finally:
        session.close()


@app.route('/api/users/me/change-password', methods=['POST'])
@require_auth
def api_force_change_password():
    user_id = request.current_user['id']
    
    # 检查速率限制
    ip = request.remote_addr
    is_allowed, remaining = check_sensitive_op_rate_limit(ip, 'change_password')
    if not is_allowed:
        return jsonify({'error': f'操作过于频繁，请{remaining}秒后重试'}), 429
    
    data, err = _get_json()
    if err:
        return err
    
    new_password = data.get('password', '')
    
    valid, msg = validate_password(new_password)
    if not valid:
        return jsonify({'error': msg}), 400
    
    # 记录操作尝试
    record_sensitive_op_attempt(ip, 'change_password')
    
    session = get_session()
    try:
        user = session.query(User).filter(User.id == user_id).first()
        if not user:
            return jsonify({'error': '用户不存在'}), 404
        
        import hashlib
        salt = secrets.token_hex(16)
        password_sha256 = hashlib.sha256(new_password.encode()).hexdigest()
        password_hash = hashlib.sha256((salt + password_sha256).encode()).hexdigest()
        user.password_hash = f"{salt}:{password_hash}"
        user.login_attempts = "0"
        user.locked_until = ""
        user.is_default = "否"
        
        session.commit()
        
        add_operation_log(
            user_id,
            request.current_user['username'],
            "修改密码",
            "",
            request.remote_addr,
            "Web"
        )
        
        return jsonify({'success': True, 'message': '密码修改成功'})
    except Exception as e:
        session.rollback()
        return jsonify({'error': str(e)}), 500
    finally:
        session.close()


@app.route('/api/statistics')
@require_auth
def api_statistics():
    session = get_session()
    try:
        total = session.query(Person).count()
        dead = session.query(Person).filter(Person.is_dead == '是').count()
        migrated = session.query(Person).filter(Person.is_migrated == '是').count()
        dibao = session.query(Person).filter(Person.dibao == '低保户').count()
        disability = session.query(Person).filter(Person.disability != '无', Person.disability != '').count()

        poor_type_distribution = {}

        hu_haos = session.query(Person.hu_hao).filter(
            Person.hu_hao != '', Person.hu_hao.isnot(None)
        ).distinct().all()

        for (hu_hao,) in hu_haos:
            members = session.query(Person.poor_type).filter(
                Person.hu_hao == hu_hao,
                Person.poor_type != '', Person.poor_type.isnot(None)
            ).all()

            if members:
                hu_zhu = session.query(Person.poor_type).filter(
                    Person.hu_hao == hu_hao,
                    Person.relation == '户主'
                ).first()

                if hu_zhu and hu_zhu[0]:
                    pt_val = hu_zhu[0]
                else:
                    pt_val = members[0][0] or '其他'

                poor_type_distribution[pt_val] = poor_type_distribution.get(pt_val, 0) + 1

        dis_level_distribution = {}
        dis_levels = session.query(Person.dis_level).filter(
            Person.dis_level != '', Person.dis_level.isnot(None), Person.dis_level != '无'
        ).all()
        for dl in dis_levels:
            dl_val = dl[0] or '其他'
            dis_level_distribution[dl_val] = dis_level_distribution.get(dl_val, 0) + 1

        land_areas = session.query(Person.land_area).filter(
            Person.land_area != '', Person.land_area.isnot(None)
        ).all()
        total_land_area = 0.0
        land_count = 0
        for la in land_areas:
            try:
                area = float(la[0])
                total_land_area += area
                land_count += 1
            except (ValueError, TypeError):
                pass
        avg_land_area = round(total_land_area / land_count, 2) if land_count > 0 else 0

        return jsonify({
            'total': total,
            'dead': dead,
            'migrated': migrated,
            'dibao': dibao,
            'disability': disability,
            'alive': total - dead,
            'not_migrated': total - migrated,
            'poor_type_distribution': poor_type_distribution,
            'dis_level_distribution': dis_level_distribution,
            'total_land_area': round(total_land_area, 2),
            'avg_land_area': avg_land_area
        })
    finally:
        session.close()


@app.route('/api/persons')
@require_auth
def api_persons():
    session = get_session()
    try:
        page = request.args.get('page', 1, type=int)
        per_page = request.args.get('per_page', 20, type=int)
        search = request.args.get('search', '')
        
        query = session.query(Person)
        if search:
            safe_search = search.replace('%', '\\%').replace('_', '\\_')
            query = query.filter(Person.name.like(f'%{safe_search}%', escape='\\'))
        
        total = query.count()
        persons = query.order_by(Person.id.desc()).limit(per_page).offset((page-1)*per_page).all()
        
        return jsonify({
            'persons': [{
                'id': p.id,
                'name': p.name,
                'id_card': p.id_card,
                'hu_hao': p.hu_hao,
                'age': p.age or '',
                'address': p.address or '',
                'phone': p.phone or '',
                'is_dead': p.is_dead or '否',
                'disability': p.disability or '',
                'dibao': p.dibao or '',
                'is_migrated': p.is_migrated or '否',
                'relation': p.relation or ''
            } for p in persons],
            'total': total,
            'page': page,
            'per_page': per_page,
            'pages': (total + per_page - 1) // per_page
        })
    finally:
        session.close()


@app.route('/api/search')
@require_auth
def api_search():
    name = request.args.get('name', '')
    if not name:
        return jsonify({'results': []})
    persons = search_by_name(name)
    
    if persons:
        add_operation_log(
            request.current_user['id'],
            request.current_user['username'],
            "搜索人员",
            f"关键词: {name}, 结果数: {len(persons)}",
            request.remote_addr,
            request.headers.get('X-Client-Type', 'Web')
        )
    
    return jsonify({'results': [{
        'id': p.id,
        'name': p.name,
        'id_card': p.id_card,
        'hu_hao': p.hu_hao,
        'age': p.age or '',
        'address': p.address or '',
        'relation': p.relation or '',
        'phone': p.phone or '',
        'is_dead': p.is_dead or '否',
        'remark': p.remark or '',
        'poor_type': p.poor_type or '',
        'disability': p.disability or '',
        'dis_level': p.dis_level or '',
        'dibao': p.dibao or '',
        'jisheng': p.jisheng or '',
        'social_card1': p.social_card1 or '',
        'social_card2': p.social_card2 or '',
        'social_card3': p.social_card3 or '',
        'is_migrated': p.is_migrated or '否',
        'land_area': p.land_area or ''
    } for p in persons]})


@app.route('/api/search/advanced')
@require_auth
def api_search_advanced():
    name = request.args.get('name', '')
    id_card = request.args.get('id_card', '')
    hu_hao = request.args.get('hu_hao', '')
    address = request.args.get('address', '')
    age_min = request.args.get('age_min', type=int)
    age_max = request.args.get('age_max', type=int)
    disability = request.args.get('disability', '')
    dibao = request.args.get('dibao', '')
    is_dead = request.args.get('is_dead', '')
    limit = request.args.get('limit', 100, type=int)
    
    persons = search_persons_advanced(
        name=name,
        id_card=id_card,
        hu_hao=hu_hao,
        address=address,
        age_min=age_min,
        age_max=age_max,
        disability=disability,
        dibao=dibao,
        is_dead=is_dead,
        limit=limit
    )
    
    return jsonify([{
        'id': p.id,
        'name': p.name,
        'id_card': p.id_card,
        'hu_hao': p.hu_hao,
        'age': p.age,
        'address': p.address,
        'relation': p.relation,
        'disability': p.disability,
        'dibao': p.dibao,
        'is_dead': p.is_dead,
        'is_migrated': p.is_migrated,
        'land_area': p.land_area or ''
    } for p in persons])


@app.route('/api/check_duplicate')
@require_auth
def api_check_duplicate():
    hu_hao = request.args.get('hu_hao', '')
    id_card = request.args.get('id_card', '')
    if not hu_hao or not id_card:
        return jsonify({'duplicate': False})
    from database import check_duplicate_id_card
    duplicate = check_duplicate_id_card(hu_hao, id_card)
    return jsonify({'duplicate': duplicate})


@app.route('/api/family/<hu_hao>')
@require_auth
def api_family(hu_hao):
    members = get_family_members(hu_hao)
    from datetime import datetime
    current_year = str(datetime.now().year)
    result = []
    for m in members:
        edu = get_student_education(m.id)
        student_mark = ""
        if edu:
            grade_display = GRADE_DISPLAY.get(edu.grade, edu.grade)
            student_mark = f"{grade_display} {edu.school_name}"
        from excel_importer import get_person_subsidy_total
        subsidy_amount = get_person_subsidy_total(m.id_card, current_year)
        result.append({
            'id': m.id,
            'name': m.name,
            'id_card': m.id_card,
            'relation': m.relation,
            'address': m.address,
            'age': m.age,
            'phone': m.phone,
            'is_dead': m.is_dead,
            'social_card1': m.social_card1,
            'social_card2': m.social_card2,
            'social_card3': m.social_card3,
            'poor_type': m.poor_type,
            'remark': m.remark,
            'disability': m.disability,
            'dis_level': m.dis_level,
            'dibao': m.dibao,
            'jisheng': m.jisheng,
            'is_migrated': m.is_migrated,
            'land_area': m.land_area or '',
            'land_area_change_reason': m.land_area_change_reason or '',
            'land_area_change_time': m.land_area_change_time or '',
            'student_mark': student_mark,
            'subsidy_amount': subsidy_amount
        })
    return jsonify({'members': result})


@app.route('/api/person/<person_id>')
@require_auth
def api_person(person_id):
    session = get_session()
    try:
        p, error = safe_id_card_query(session, person_id)
        if error:
            return error
        
        add_operation_log(
            request.current_user['id'],
            request.current_user['username'],
            "查看人员详情",
            f"{p.name}({p.id_card[:6]}****{p.id_card[-4:]})" if p.id_card else p.name,
            request.remote_addr,
            request.headers.get('X-Client-Type', 'Web')
        )
        
        edu = get_student_education(p.id)
        return jsonify({
            'person': {
                'id': p.id,
                'name': p.name,
                'id_card': p.id_card,
                'hu_hao': p.hu_hao,
                'relation': p.relation,
                'address': p.address,
                'age': p.age,
                'phone': p.phone,
                'is_dead': p.is_dead,
                'social_card1': p.social_card1,
                'social_card2': p.social_card2,
                'social_card3': p.social_card3,
                'poor_type': p.poor_type,
                'remark': p.remark,
                'disability': p.disability,
                'disability_type': p.disability_type,
                'dis_level': p.dis_level,
                'dibao': p.dibao,
                'jisheng': p.jisheng,
                'is_migrated': p.is_migrated,
                'land_area': p.land_area or '',
                'land_area_change_reason': p.land_area_change_reason or '',
                'land_area_change_time': p.land_area_change_time or '',
                'updated_at': p.updated_at or '',
                'education': {
                    'school_name': edu.school_name,
                    'grade': edu.grade,
                    'class_name': edu.class_name,
                    'guardian': edu.guardian,
                    'guardian_phone': edu.guardian_phone,
                    'is_left_behind': edu.is_left_behind,
                    'student_status': edu.student_status,
                    'enrollment_date': edu.enrollment_date,
                    'graduation_date': edu.graduation_date
                } if edu else None
            }
        })
    finally:
        session.close()


@app.route('/api/person', methods=['POST'])
@require_auth
def api_add_person():
    if request.current_user.get('role') not in ['admin', 'user']:
        return jsonify({'error': '权限不足'}), 403
    data, err = _get_json()
    if err:
        return err

    hu_hao = data.get('hu_hao', '').strip()
    if not hu_hao:
        hu_hao = f"H{secrets.token_hex(4).upper()}"
        data['hu_hao'] = hu_hao

    person = Person(
        id=f"web_{secrets.token_hex(8)}",
        hu_hao=hu_hao,
        id_card=data['id_card'],
        name=data['name'],
        age=calculate_age(data['id_card']),
        address=data.get('address', ''),
        relation=data.get('relation', ''),
        phone=data.get('phone', '无'),
        is_dead=data.get('is_dead', '否'),
        land_area=data.get('land_area', ''),
        social_card1=data.get('social_card1', ''),
        social_card2=data.get('social_card2', ''),
        social_card3=data.get('social_card3', ''),
        poor_type=data.get('poor_type', ''),
        remark=data.get('remark', ''),
        disability=data.get('disability', ''),
        disability_type=data.get('disability_type', ''),
        dis_level=data.get('dis_level', ''),
        dibao=data.get('dibao', ''),
        jisheng=data.get('jisheng', ''),
    )
    add_person(person)
    add_operation_log(request.current_user['id'], request.current_user['username'], "添加人员", data.get('name', ''), request.remote_addr, "Web")
    return jsonify({'success': True})


@app.route('/api/person/<person_id>', methods=['PUT'])
@require_auth
def api_update_person(person_id):
    if request.current_user.get('role') not in ['admin', 'user']:
        return jsonify({'error': '权限不足'}), 403
    data, err = _get_json()
    if err:
        return err
    client_updated_at = data.get('updated_at', '')
    session = get_session()
    try:
        p, error = safe_id_card_query(session, person_id)
        if error:
            return error
        if client_updated_at and p.updated_at and client_updated_at != p.updated_at:
            return jsonify({
                'error': 'conflict',
                'message': f'数据已被其他人修改！\n上次修改时间：{p.updated_at}\n您的修改时间：{client_updated_at}\n请刷新后重试。'
            }), 409
        
        if 'id_card' in data and data['id_card'] != p.id_card:
            return jsonify({'error': '身份证号不允许修改，如需更正请删除后重新添加'}), 400
        
        p.name = data.get('name', p.name)
        p.relation = data.get('relation', p.relation)
        p.phone = data.get('phone', p.phone)
        p.is_dead = data.get('is_dead', p.is_dead)
        p.address = data.get('address', p.address)
        p.poor_type = data.get('poor_type', p.poor_type)
        p.remark = data.get('remark', p.remark)
        p.disability = data.get('disability', p.disability)
        p.disability_type = data.get('disability_type', p.disability_type)
        p.dis_level = data.get('dis_level', p.dis_level)
        p.dibao = data.get('dibao', p.dibao)
        p.jisheng = data.get('jisheng', p.jisheng)
        p.social_card1 = data.get('social_card1', p.social_card1)
        p.social_card2 = data.get('social_card2', p.social_card2)
        p.social_card3 = data.get('social_card3', p.social_card3)
        new_land_area = data.get('land_area', p.land_area)
        old_land_area = p.land_area or ''
        
        if new_land_area != old_land_area and old_land_area:
            land_reason = data.get('land_area_change_reason', '')
            if not land_reason:
                return jsonify({'error': '土地面积变更时必须填写变更原因'}), 400
            p.land_area = new_land_area
            p.land_area_change_reason = land_reason
            p.land_area_change_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        elif new_land_area != old_land_area:
            p.land_area = new_land_area
        update_person(p)
        add_operation_log(request.current_user['id'], request.current_user['username'], "修改人员", p.name, request.remote_addr, "Web")
        return jsonify({'success': True, 'updated_at': p.updated_at})
    finally:
        session.close()


@app.route('/api/person/<person_id>', methods=['DELETE'])
@require_auth
def api_delete_person(person_id):
    if request.current_user.get('role') != 'admin':
        return jsonify({'error': '权限不足'}), 403
    session = get_session()
    try:
        p, error = safe_id_card_query(session, person_id)
        if error:
            return error
        real_id = p.id
    finally:
        session.close()
    delete_person(real_id)
    add_operation_log(request.current_user['id'], request.current_user['username'], "删除人员", person_id, request.remote_addr, "Web", "是")
    return jsonify({'success': True})


@app.route('/api/person/<person_id>/migrate', methods=['POST'])
@require_auth
def api_migrate_person(person_id):
    if request.current_user.get('role') not in ['admin', 'user']:
        return jsonify({'error': '权限不足'}), 403
    data, err = _get_json()
    if err:
        return err
    session = get_session()
    try:
        p, error = safe_id_card_query(session, person_id)
        if error:
            return error
        real_id = p.id
    finally:
        session.close()
    migrate_person(
        real_id,
        data.get('date', ''),
        data.get('place', ''),
        data.get('reason', '')
    )
    add_operation_log(request.current_user['id'], request.current_user['username'], "户籍迁出", person_id, request.remote_addr, "Web")
    return jsonify({'success': True})


@app.route('/api/household/split', methods=['POST'])
@require_auth
def api_split_household():
    if request.current_user.get('role') not in ['admin', 'user']:
        return jsonify({'error': '权限不足'}), 403
    data, err = _get_json()
    if err:
        return err
    person_ids = data.get('person_ids', [])
    new_hu_hao = data.get('new_hu_hao', '')
    from database import split_household
    try:
        result_hu_hao = split_household(person_ids, new_hu_hao)
        add_operation_log(request.current_user['id'], request.current_user['username'], "分户", f"新户号:{result_hu_hao}", request.remote_addr, "Web")
        return jsonify({'success': True, 'new_hu_hao': result_hu_hao})
    except Exception as e:
        return jsonify({'error': str(e)}), 400


@app.route('/api/household/merge', methods=['POST'])
@require_auth
def api_merge_household():
    if request.current_user.get('role') not in ['admin', 'user']:
        return jsonify({'error': '权限不足'}), 403
    data, err = _get_json()
    if err:
        return err
    target_hu_hao = data.get('target_hu_hao', '')
    person_ids = data.get('person_ids', [])
    from database import merge_household
    try:
        merge_household(target_hu_hao, person_ids)
        add_operation_log(request.current_user['id'], request.current_user['username'], "合户", f"目标户号:{target_hu_hao}", request.remote_addr, "Web")
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 400


@app.route('/api/household/add', methods=['POST'])
@require_auth
def api_add_household():
    if request.current_user.get('role') not in ['admin', 'user']:
        return jsonify({'error': '权限不足'}), 403
    data, err = _get_json()
    if err:
        return err
    from database import add_new_household
    try:
        add_new_household(
            hu_hao=data.get('hu_hao', ''),
            name=data.get('name', ''),
            id_card=data.get('id_card', ''),
            relation=data.get('relation', '户主'),
            address=data.get('address', ''),
            phone=data.get('phone', ''),
            land_area=data.get('land_area', '')
        )
        add_operation_log(request.current_user['id'], request.current_user['username'], "新增户", f"户号:{data.get('hu_hao')}", request.remote_addr, "Web")
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 400


@app.route('/api/household/new_number')
@require_auth
def api_new_hu_hao():
    from database import generate_new_hu_hao
    return jsonify({'hu_hao': generate_new_hu_hao()})


@app.route('/api/students')
@require_auth
def api_students():
    school = request.args.get('school')
    grade = request.args.get('grade')
    class_name = request.args.get('class_name')
    is_left_behind = request.args.get('is_left_behind')
    students = search_students(school, grade, class_name, is_left_behind)
    session = get_session()
    try:
        result = []
        for s in students:
            p = session.query(Person).filter(Person.id == s.person_id).first()
            result.append({
                'id': s.id,
                'person_id': s.person_id,
                'name': p.name if p else '',
                'id_card': s.id_card,
                'school_name': s.school_name,
                'grade': s.grade,
                'grade_display': GRADE_DISPLAY.get(s.grade, s.grade),
                'class_name': s.class_name,
                'guardian': s.guardian,
                'is_left_behind': s.is_left_behind,
                'student_status': s.student_status
            })
        return jsonify({'students': result})
    finally:
        session.close()


@app.route('/api/stats')
@require_auth
def api_stats():
    return jsonify({
        'total_persons': get_person_count()
    })


@app.route('/api/statistics/trend')
@require_auth
def api_statistics_trend():
    start_date = request.args.get('start_date', '')
    end_date = request.args.get('end_date', '')
    group_by = request.args.get('group_by', 'month')
    
    if group_by not in ['day', 'week', 'month', 'year']:
        group_by = 'month'
    
    if not start_date or not end_date:
        from datetime import datetime, timedelta
        end_dt = datetime.now()
        start_dt = end_dt - timedelta(days=365)
        start_date = start_dt.strftime('%Y-%m-%d')
        end_date = end_dt.strftime('%Y-%m-%d')
    
    try:
        data = get_person_trend_data(start_date, end_date, group_by)
        return jsonify(data)
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/api/statistics/distribution')
@require_auth
def api_statistics_distribution():
    field = request.args.get('field', '')
    
    if not field:
        return jsonify({'error': '缺少field参数'}), 400
    
    valid_fields = ['disability', 'dibao', 'age_group', 'gender']
    if field not in valid_fields and not hasattr(Person, field):
        return jsonify({'error': f'无效的field参数，支持: {", ".join(valid_fields)}'}), 400
    
    try:
        data = get_person_distribution(field)
        return jsonify(data)
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/api/statistics/comparison')
@require_auth
def api_statistics_comparison():
    dimensions = request.args.getlist('dimensions')
    
    if not dimensions:
        dimensions = ['age', 'disability', 'dibao', 'migration', 'student']
    
    valid_dims = ['age', 'disability', 'dibao', 'migration', 'death', 'student', 'gender_male', 'gender_female']
    
    try:
        data = get_person_comparison(dimensions)
        return jsonify(data)
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/api/students/statistics')
@require_auth
def api_students_statistics():
    session = get_session()
    try:
        total = session.query(StudentEducation).count()
        
        studying = session.query(StudentEducation).filter(
            StudentEducation.student_status == '在读'
        ).count()
        
        graduated = session.query(StudentEducation).filter(
            StudentEducation.student_status == '离校'
        ).count()
        
        left_behind = session.query(StudentEducation).filter(
            StudentEducation.is_left_behind == '是'
        ).count()
        
        grade_stats = {}
        grades = session.query(StudentEducation.grade, func.count(StudentEducation.id)).group_by(
            StudentEducation.grade
        ).all()
        for grade, count in grades:
            if grade:
                grade_stats[grade] = count
        
        school_stats = {}
        schools = session.query(StudentEducation.school_name, func.count(StudentEducation.id)).group_by(
            StudentEducation.school_name
        ).all()
        for school, count in schools:
            if school:
                school_stats[school] = count
        
        return jsonify({
            'total': total,
            'studying': studying,
            'graduated': graduated,
            'left_behind': left_behind,
            'grade_distribution': grade_stats,
            'school_distribution': school_stats
        })
    finally:
        session.close()


@app.route('/api/statistics/age_distribution')
@require_auth
def api_statistics_age_distribution():
    session = get_session()
    try:
        age_ranges = [
            ('0-18岁', 0, 18),
            ('19-30岁', 19, 30),
            ('31-40岁', 31, 40),
            ('41-50岁', 41, 50),
            ('51-60岁', 51, 60),
            ('60岁以上', 61, 200)
        ]
        
        distribution = []
        for label, min_age, max_age in age_ranges:
            count = session.query(Person).filter(
                Person.age != '',
                Person.age != None
            ).all()
            
            count_num = 0
            for p in count:
                try:
                    age = int(p.age)
                    if min_age <= age <= max_age:
                        count_num += 1
                except Exception:
                    pass
            
            distribution.append({
                'range': label,
                'count': count_num
            })
        
        return jsonify({'distribution': distribution})
    finally:
        session.close()


@app.route('/api/statistics/by_village')
@require_auth
def api_statistics_by_village():
    session = get_session()
    try:
        villages = session.query(Person.address, func.count(Person.id)).group_by(
            Person.address
        ).all()
        
        village_stats = []
        for address, count in villages:
            if address:
                village_stats.append({
                    'village': address,
                    'count': count
                })
        
        return jsonify({'villages': village_stats[:50]})
    finally:
        session.close()


PHOTO_UPLOAD_DIR = os.path.join(BASE_DIR, 'Photos')
os.makedirs(PHOTO_UPLOAD_DIR, exist_ok=True)


@app.route('/api/upload_photo', methods=['POST'])
@require_auth
def api_upload_photo_by_id_card():
    if request.current_user.get('role') not in ['admin', 'user']:
        return jsonify({'error': '权限不足'}), 403
    
    id_card = request.form.get('id_card', '')
    photo_type = request.form.get('photo_type', '身份证')
    
    if not id_card:
        return jsonify({'error': '身份证号不能为空'}), 400
    
    if 'photo' not in request.files and 'file' not in request.files:
        return jsonify({'error': '未选择文件'}), 400
    
    file = request.files.get('photo') or request.files.get('file')
    if not file or not file.filename:
        return jsonify({'error': '未选择文件'}), 400
    
    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in PHOTO_EXTENSIONS:
        return jsonify({'error': f'不支持的文件格式，仅支持：{", ".join(PHOTO_EXTENSIONS)}'}), 400
    
    img_type = validate_image_content(file)
    if not img_type:
        return jsonify({'error': '文件内容不是有效的图片格式'}), 400
    
    session = get_session()
    try:
        p = session.query(Person).filter(Person.id_card == id_card).first()
        if not p:
            return jsonify({'error': '人员不存在'}), 404
        
        base_name = get_photo_base_name(photo_type, p.id_card, p.social_card1, p.social_card2, p.social_card3)
        if not base_name:
            return jsonify({'error': '无法生成照片文件名'}), 400
        
        cleaned = clean_filename(base_name)
        old = find_existing_photo(base_name, PHOTO_UPLOAD_DIR)
        if old and os.path.exists(old):
            os.remove(old)
        
        dest = os.path.join(PHOTO_UPLOAD_DIR, cleaned + ext)
        file.save(dest)
        
        photo_url = f'/api/photo/{cleaned}{ext}'
        
        return jsonify({
            'success': True,
            'url': photo_url,
            'message': '上传成功'
        })
    finally:
        session.close()


@app.route('/api/person/<person_id>/photo', methods=['POST'])
@require_auth
def api_upload_photo(person_id):
    if request.current_user.get('role') not in ['admin', 'user']:
        return jsonify({'error': '权限不足'}), 403

    session = get_session()
    try:
        p, error = safe_id_card_query(session, person_id)
        if error:
            return error

        photo_type = request.form.get('photo_type', '身份证')
        if 'file' not in request.files:
            return jsonify({'error': '未选择文件'}), 400

        file = request.files['file']
        if not file.filename:
            return jsonify({'error': '未选择文件'}), 400

        ext = os.path.splitext(file.filename)[1].lower()
        if ext not in PHOTO_EXTENSIONS:
            return jsonify({'error': f'不支持的文件格式，仅支持：{", ".join(PHOTO_EXTENSIONS)}'}), 400

        img_type = validate_image_content(file)
        if not img_type:
            return jsonify({'error': '文件内容不是有效的图片格式'}), 400

        base_name = get_photo_base_name(photo_type, p.id_card, p.social_card1, p.social_card2, p.social_card3)
        if not base_name:
            return jsonify({'error': '无法生成照片文件名'}), 400

        cleaned = clean_filename(base_name)
        old = find_existing_photo(base_name, PHOTO_UPLOAD_DIR)
        if old and os.path.exists(old):
            os.remove(old)

        dest = os.path.join(PHOTO_UPLOAD_DIR, cleaned + ext)
        file.save(dest)

        photo_url = f'/api/person/{p.id_card}/photo'

        return jsonify({
            'success': True,
            'url': photo_url,
            'message': '上传成功'
        })
    finally:
        session.close()


@app.route('/api/photo/<id_card>/<photo_type>')
@require_auth
def api_get_photo_by_type(id_card, photo_type):
    session = get_session()
    try:
        p = session.query(Person).filter(Person.id_card == id_card).first()
        if not p:
            return jsonify({'error': '人员不存在'}), 404
        
        base_name = get_photo_base_name(photo_type, p.id_card, p.social_card1, p.social_card2, p.social_card3)
        if not base_name:
            return jsonify({'error': '无法定位照片'}), 400
        
        path = find_existing_photo(base_name, PHOTO_UPLOAD_DIR)
        if path and os.path.exists(path):
            return send_file(path, mimetype='image/jpeg')
        
        return jsonify({'error': '照片不存在'}), 404
    finally:
        session.close()


@app.route('/api/photo/<filename>')
@require_auth
def api_serve_photo(filename):
    try:
        file_path = os.path.join(PHOTO_UPLOAD_DIR, filename)
        if not os.path.exists(file_path):
            return jsonify({'error': '照片不存在'}), 404
        
        real_path = os.path.realpath(file_path)
        if not real_path.startswith(os.path.realpath(PHOTO_UPLOAD_DIR)):
            return jsonify({'error': '无权访问'}), 403
        
        return send_file(file_path)
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/api/person/<person_id>/photo')
@require_auth
def api_get_photo(person_id):
    session = get_session()
    try:
        p, error = safe_id_card_query(session, person_id)
        if error:
            return error

        base_name = get_photo_base_name('身份证', p.id_card, p.social_card1, p.social_card2, p.social_card3)
        if not base_name:
            return jsonify({'error': '无法查找照片'}), 400

        path = find_existing_photo(base_name, PHOTO_UPLOAD_DIR)
        if path and os.path.exists(path):
            return send_file(path, mimetype='image/jpeg')

        for card in [p.social_card1, p.social_card2, p.social_card3]:
            if card and card != "无":
                base = get_photo_base_name('身份证', card, '', '', '')
                if base:
                    path = find_existing_photo(base, PHOTO_UPLOAD_DIR)
                    if path and os.path.exists(path):
                        return send_file(path, mimetype='image/jpeg')

        return jsonify({'error': '照片不存在'}), 404
    finally:
        session.close()


@app.route('/api/person/<person_id>/photo', methods=['DELETE'])
@require_auth
def api_delete_photo(person_id):
    if request.current_user.get('role') not in ['admin', 'user']:
        return jsonify({'error': '权限不足'}), 403

    session = get_session()
    try:
        p, error = safe_id_card_query(session, person_id)
        if error:
            return error

        photo_type = request.args.get('photo_type', '身份证')
        base_name = get_photo_base_name(photo_type, p.id_card, p.social_card1, p.social_card2, p.social_card3)
        if not base_name:
            return jsonify({'error': '无法生成照片文件名'}), 400

        path = find_existing_photo(base_name, PHOTO_UPLOAD_DIR)
        if path and os.path.exists(path):
            os.remove(path)
            return jsonify({'success': True})

        return jsonify({'error': '照片不存在'}), 404
    finally:
        session.close()


@app.route('/api/import/excel', methods=['POST'])
@require_auth
def api_import_excel():
    if request.current_user.get('role') != 'admin':
        return jsonify({'error': '权限不足，仅管理员可导入'}), 403
    if 'file' not in request.files:
        return jsonify({'error': '未选择文件'}), 400
    file = request.files['file']
    if not file.filename:
        return jsonify({'error': '未选择文件'}), 400
    mode = request.form.get('mode', 'incremental')
    tmp_path = os.path.join(BASE_DIR, f'_tmp_import_{secrets.token_hex(4)}.xlsx')
    try:
        file.save(tmp_path)
        from excel_importer import import_excel_full, import_excel_incremental
        if mode == 'full':
            result = import_excel_full(tmp_path)
        else:
            result = import_excel_incremental(tmp_path)
        add_operation_log(request.current_user['id'], request.current_user['username'], "导入Excel", f"模式:{mode} 成功:{result.get('imported',0)}", request.remote_addr, "Web")
        return jsonify(result)
    except Exception as e:
        return jsonify({'error': str(e)}), 400
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


@app.route('/api/export/excel', methods=['POST'])
@require_auth
def api_export_excel():
    data = request.get_json() or {}
    password = data.get('password', '')
    
    if not password:
        return jsonify({'error': '导出敏感数据需要验证密码'}), 401
    
    user_id = request.current_user['id']
    session = get_session()
    try:
        user = session.query(User).filter(User.id == user_id).first()
        if not user:
            return jsonify({'error': '用户不存在'}), 404
        
        import hashlib
        salt, stored_hash = user.password_hash.split(":")
        password_sha256 = hashlib.sha256(password.encode()).hexdigest()
        password_hash = hashlib.sha256((salt + password_sha256).encode()).hexdigest()
        if password_hash != stored_hash:
            old_hash = hashlib.sha256((salt + password).encode()).hexdigest()
            if old_hash != stored_hash:
                return jsonify({'error': '密码错误'}), 401
    finally:
        session.close()
    
    from excel_importer import export_excel
    filters = data.get('filters')
    tmp_path = os.path.join(BASE_DIR, f'_tmp_export_{secrets.token_hex(4)}.xlsx')
    try:
        if export_excel(tmp_path, filters=filters):
            add_operation_log(
                request.current_user['id'],
                request.current_user['username'],
                "导出敏感数据",
                f"导出人口信息Excel，筛选条件: {json.dumps(filters) if filters else '无'}",
                request.remote_addr,
                "Web"
            )
            return send_file(tmp_path, as_attachment=True, download_name=f'人口信息_{datetime.now().strftime("%Y%m%d")}.xlsx')
        return jsonify({'error': '没有数据可导出'}), 404
    except Exception as e:
        return jsonify({'error': str(e)}), 400


@app.route('/api/template/excel')
@require_auth
def api_download_template():
    from excel_importer import generate_template
    tmp_path = os.path.join(BASE_DIR, f'_tmp_template_{secrets.token_hex(4)}.xlsx')
    try:
        generate_template(tmp_path)
        return send_file(tmp_path, as_attachment=True, download_name='人口信息导入模板.xlsx')
    except Exception as e:
        return jsonify({'error': str(e)}), 400


@app.route('/api/template/student')
@require_auth
def api_download_student_template():
    from student_importer import export_student_education_template
    tmp_path = os.path.join(BASE_DIR, f'_tmp_stu_template_{secrets.token_hex(4)}.xlsx')
    try:
        export_student_education_template(tmp_path)
        return send_file(tmp_path, as_attachment=True, download_name='学生就读信息模板.xlsx')
    except Exception as e:
        return jsonify({'error': str(e)}), 400


@app.route('/api/deleted')
@require_auth
def api_deleted_persons():
    if request.current_user.get('role') != 'admin':
        return jsonify({'error': '权限不足'}), 403
    from database import get_deleted_persons
    persons = get_deleted_persons()
    return jsonify([{
        'id': p.id,
        'original_id': p.original_id,
        'name': p.name,
        'id_card': p.id_card,
        'hu_hao': p.hu_hao,
        'deleted_at': p.deleted_at
    } for p in persons])


@app.route('/api/recover/<deleted_id>', methods=['POST'])
@require_auth
def api_recover_person(deleted_id):
    if request.current_user.get('role') != 'admin':
        return jsonify({'error': '权限不足'}), 403
    from database import recover_person
    try:
        recover_person(deleted_id)
        add_operation_log(request.current_user['id'], request.current_user['username'], "找回人员", deleted_id, request.remote_addr, "Web")
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 400


@app.route('/api/students/search')
@require_auth
def api_students_search():
    name = request.args.get('name')
    school = request.args.get('school')
    grade = request.args.get('grade')
    class_name = request.args.get('class_name')
    is_left_behind = request.args.get('is_left_behind')

    session = get_session()
    try:
        query = session.query(StudentEducation).join(
            Person, StudentEducation.person_id == Person.id
        )

        if name and name.strip():
            safe_name = name.strip().replace('%', '\\%').replace('_', '\\_')
            query = query.filter(Person.name.like(f'%{safe_name}%', escape='\\'))

        if school and school.strip():
            safe_school = school.strip().replace('%', '\\%').replace('_', '\\_')
            query = query.filter(StudentEducation.school_name.like(f'%{safe_school}%', escape='\\'))

        if grade and grade.strip():
            query = query.filter(StudentEducation.grade == grade.strip())

        if class_name and class_name.strip():
            safe_class = class_name.strip().replace('%', '\\%').replace('_', '\\_')
            query = query.filter(StudentEducation.class_name.like(f'%{safe_class}%', escape='\\'))

        if is_left_behind and is_left_behind.strip():
            query = query.filter(StudentEducation.is_left_behind == is_left_behind.strip())

        students = query.limit(100).all()

        result = []
        for s in students:
            p = session.query(Person).filter(Person.id == s.person_id).first()
            result.append({
                'id': s.id,
                'person_id': s.person_id,
                'name': p.name if p else '',
                'id_card': s.id_card,
                'school_name': s.school_name,
                'grade': s.grade,
                'grade_display': GRADE_DISPLAY.get(s.grade, s.grade),
                'class_name': s.class_name,
                'guardian': s.guardian,
                'guardian_phone': s.guardian_phone,
                'is_left_behind': s.is_left_behind,
                'student_status': s.student_status,
                'enrollment_date': s.enrollment_date,
                'graduation_date': s.graduation_date
            })
        return jsonify({'results': result})
    finally:
        session.close()


@app.route('/api/person/<person_id>/education', methods=['PUT'])
@require_auth
def api_save_education(person_id):
    if request.current_user.get('role') not in ['admin', 'user']:
        return jsonify({'error': '权限不足'}), 403
    data, err = _get_json()
    if err:
        return err
    session = get_session()
    try:
        p, error = safe_id_card_query(session, person_id)
        if error:
            return error

        from database import StudentEducation, save_student_education, delete_student_education
        edu = session.query(StudentEducation).filter(StudentEducation.person_id == p.id).first()

        school = data.get('school_name', '').strip()
        grade = data.get('grade', '')
        class_name = data.get('class_name', '').strip()

        if not school and not grade and not class_name:
            if edu:
                delete_student_education(p.id)
            return jsonify({'success': True})

        if edu:
            edu.school_name = school
            edu.grade = grade
            edu.class_name = class_name
            edu.guardian = data.get('guardian', '')
            edu.guardian_phone = data.get('guardian_phone', '')
            edu.is_left_behind = data.get('is_left_behind', '否')
            edu.student_status = data.get('student_status', '在读')
            edu.enrollment_date = data.get('enrollment_date', '')
            edu.graduation_date = data.get('graduation_date', '')
        else:
            edu = StudentEducation(
                id=f"edu_{p.id}",
                person_id=p.id,
                id_card=p.id_card,
                school_name=school,
                grade=grade,
                class_name=class_name,
                guardian=data.get('guardian', ''),
                guardian_phone=data.get('guardian_phone', ''),
                is_left_behind=data.get('is_left_behind', '否'),
                student_status=data.get('student_status', '在读'),
                enrollment_date=data.get('enrollment_date', ''),
                graduation_date=data.get('graduation_date', '')
            )
            session.add(edu)
        session.commit()
        add_operation_log(request.current_user['id'], request.current_user['username'], "修改就读信息", p.name, request.remote_addr, "Web")
        return jsonify({'success': True})
    except Exception as e:
        session.rollback()
        return jsonify({'error': str(e)}), 400
    finally:
        session.close()


@app.route('/api/students/import', methods=['POST'])
@require_auth
def api_students_import():
    if request.current_user.get('role') not in ['admin', 'user']:
        return jsonify({'error': '权限不足'}), 403
    if 'file' not in request.files:
        return jsonify({'error': '未选择文件'}), 400
    file = request.files['file']
    if not file.filename:
        return jsonify({'error': '未选择文件'}), 400
    tmp_path = os.path.join(BASE_DIR, f'_tmp_stu_{secrets.token_hex(4)}.xlsx')
    try:
        file.save(tmp_path)
        from student_importer import import_student_education
        result = import_student_education(tmp_path)
        return jsonify(result)
    except Exception as e:
        return jsonify({'error': str(e)}), 400
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


@app.route('/api/students/export', methods=['POST'])
@require_auth
def api_students_export():
    if request.current_user.get('role') not in ['admin', 'user']:
        return jsonify({'error': '权限不足'}), 403
    
    data = request.get_json() or {}
    password = data.get('password', '')
    
    if not password:
        return jsonify({'error': '导出敏感数据需要验证密码'}), 401
    
    user_id = request.current_user['id']
    session = get_session()
    try:
        user = session.query(User).filter(User.id == user_id).first()
        if not user:
            return jsonify({'error': '用户不存在'}), 404
        
        import hashlib
        salt, stored_hash = user.password_hash.split(":")
        password_sha256 = hashlib.sha256(password.encode()).hexdigest()
        password_hash = hashlib.sha256((salt + password_sha256).encode()).hexdigest()
        if password_hash != stored_hash:
            old_hash = hashlib.sha256((salt + password).encode()).hexdigest()
            if old_hash != stored_hash:
                return jsonify({'error': '密码错误'}), 401
    finally:
        session.close()
    
    from student_importer import export_student_education
    tmp_path = os.path.join(BASE_DIR, f'_tmp_stu_exp_{secrets.token_hex(4)}.xlsx')
    try:
        if export_student_education(tmp_path):
            add_operation_log(
                request.current_user['id'],
                request.current_user['username'],
                "导出敏感数据",
                "导出学生就读信息Excel",
                request.remote_addr,
                request.headers.get('X-Client-Type', 'Web')
            )
            return send_file(tmp_path, as_attachment=True, download_name=f'学生就读信息_{datetime.now().strftime("%Y%m%d")}.xlsx')
        return jsonify({'error': '没有数据可导出'}), 404
    except Exception as e:
        return jsonify({'error': str(e)}), 400


@app.route('/api/students/promote', methods=['POST'])
@require_auth
def api_promote_grades():
    if request.current_user.get('role') != 'admin':
        return jsonify({'error': '权限不足'}), 403
    from database import promote_all_grades
    try:
        count = promote_all_grades()
        add_operation_log(request.current_user['id'], request.current_user['username'], "年级升级",
                         f"升级{count}名学生", request.remote_addr, "Web")
        return jsonify({'success': True, 'count': count})
    except Exception as e:
        return jsonify({'error': str(e)}), 400


@app.route('/api/land_subsidy/import', methods=['POST'])
@require_auth
def api_import_land_subsidy():
    if request.current_user.get('role') != 'admin':
        return jsonify({'error': '权限不足'}), 403

    subsidy_year = request.form.get('subsidy_year', '')
    if 'file' not in request.files:
        return jsonify({'error': '未选择文件'}), 400

    file = request.files['file']
    if not file.filename:
        return jsonify({'error': '未选择文件'}), 400

    import tempfile
    with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
        file.save(tmp.name)
        tmp_path = tmp.name

    try:
        from excel_importer import import_land_subsidy
        result = import_land_subsidy(tmp_path, subsidy_year)
        add_operation_log(request.current_user['id'], request.current_user['username'],
                         "导入地力补贴", f"年度:{subsidy_year}, 导入{result.get('imported', 0)}条",
                         request.remote_addr, "Web")
        return jsonify(result)
    except Exception as e:
        return jsonify({'error': str(e)}), 400
    finally:
        os.remove(tmp_path)


@app.route('/api/land_subsidy')
@require_auth
def api_get_land_subsidies():
    subsidy_year = request.args.get('year', None)
    search = request.args.get('search', None)

    from excel_importer import get_all_land_subsidies
    subsidies = get_all_land_subsidies(subsidy_year)

    if search:
        subsidies = [s for s in subsidies if search in s.name]

    result = []
    for s in subsidies:
        result.append({
            'id': s.id,
            'name': s.name,
            'id_card': s.id_card,
            'subsidy_year': s.subsidy_year,
            'certified_area': s.certified_area,
            'subsidy_area': s.subsidy_area,
            'base_standard': s.base_standard,
            'rape_area': s.rape_area,
            'rape_amount': s.rape_amount,
            'total_amount': s.total_amount,
            'import_time': s.import_time,
        })

    return jsonify({'subsidies': result})


@app.route('/api/land_subsidy/years')
@require_auth
def api_get_land_subsidy_years():
    from excel_importer import get_all_land_subsidies
    all_data = get_all_land_subsidies()
    years = list(set(s.subsidy_year for s in all_data if s.subsidy_year))
    years.sort(reverse=True)
    return jsonify({'years': years})


@app.route('/api/land_subsidy/statistics')
@require_auth
def api_get_land_subsidy_statistics():
    subsidy_year = request.args.get('year', None)

    from excel_importer import get_all_land_subsidies
    all_data = get_all_land_subsidies(subsidy_year)

    if not all_data:
        return jsonify({
            'total_count': 0,
            'total_certified_area': 0,
            'total_subsidy_area': 0,
            'total_amount': 0,
            'avg_subsidy_amount': 0,
            'year_distribution': {}
        })

    total_count = len(all_data)
    total_certified_area = 0.0
    total_subsidy_area = 0.0
    total_amount = 0.0
    year_distribution = {}

    for s in all_data:
        try:
            total_certified_area += float(s.certified_area or 0)
        except (ValueError, TypeError):
            pass
        try:
            total_subsidy_area += float(s.subsidy_area or 0)
        except (ValueError, TypeError):
            pass
        try:
            total_amount += float(s.total_amount or 0)
        except (ValueError, TypeError):
            pass

        if s.subsidy_year:
            year_distribution[s.subsidy_year] = year_distribution.get(s.subsidy_year, 0) + 1

    avg_subsidy_amount = round(total_amount / total_count, 2) if total_count > 0 else 0

    return jsonify({
        'total_count': total_count,
        'total_certified_area': round(total_certified_area, 2),
        'total_subsidy_area': round(total_subsidy_area, 2),
        'total_amount': round(total_amount, 2),
        'avg_subsidy_amount': avg_subsidy_amount,
        'year_distribution': year_distribution
    })


@app.route('/api/land_subsidy/<subsidy_id>', methods=['DELETE'])
@require_auth
def api_delete_land_subsidy(subsidy_id):
    if request.current_user.get('role') != 'admin':
        return jsonify({'error': '权限不足'}), 403

    from excel_importer import delete_land_subsidy
    try:
        delete_land_subsidy(subsidy_id)
        add_operation_log(request.current_user['id'], request.current_user['username'],
                         "删除地力补贴记录", f"ID:{subsidy_id}",
                         request.remote_addr, "Web")
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 400


@app.route('/api/land_subsidy/export')
@require_auth
def api_export_land_subsidy():
    subsidy_year = request.args.get('year', None)
    search = request.args.get('search', None)

    from excel_importer import get_all_land_subsidies
    subsidies = get_all_land_subsidies(subsidy_year)

    if search:
        subsidies = [s for s in subsidies if search in s.name]

    import pandas as pd
    import io
    from flask import Response
    data = []
    for s in subsidies:
        data.append({
            '补贴年度': s.subsidy_year,
            '姓名': s.name,
            '身份证号': s.id_card,
            '确权颁发面积': s.certified_area,
            '补贴面积(亩)': s.subsidy_area,
            '基础补贴标准(元/亩)': s.base_standard,
            '油菜面积(亩)': s.rape_area,
            '油菜金额(元)': s.rape_amount,
            '补贴总金额(元)': s.total_amount,
            '导入时间': s.import_time,
        })

    if not data:
        return jsonify({'error': '没有数据可导出'}), 400

    df = pd.DataFrame(data)
    output = io.BytesIO()
    df.to_excel(output, index=False, engine='openpyxl')
    output.seek(0)

    from datetime import datetime
    filename = f"地力补贴_{datetime.now().strftime('%Y%m%d')}.xlsx"

    add_operation_log(request.current_user['id'], request.current_user['username'],
                     "导出地力补贴", f"导出{len(data)}条记录",
                     request.remote_addr, "Web")

    return Response(
        output.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename="{filename}"'}
    )


@app.route('/api/treatment_payment/import', methods=['POST'])
@require_auth
def api_import_treatment_payment():
    if request.current_user.get('role') != 'admin':
        return jsonify({'error': '权限不足'}), 403

    if 'file' not in request.files:
        return jsonify({'error': '未选择文件'}), 400

    file = request.files['file']
    if not file.filename:
        return jsonify({'error': '未选择文件'}), 400

    import tempfile
    with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
        file.save(tmp.name)
        tmp_path = tmp.name

    try:
        from excel_importer import import_treatment_payment
        result = import_treatment_payment(tmp_path)
        add_operation_log(request.current_user['id'], request.current_user['username'],
                         "导入待遇发放", f"导入{result.get('imported', 0)}条",
                         request.remote_addr, "Web")
        return jsonify(result)
    except Exception as e:
        return jsonify({'error': str(e)}), 400
    finally:
        os.remove(tmp_path)


@app.route('/api/treatment_payment')
@require_auth
def api_get_treatment_payments():
    search = request.args.get('search', None)

    from excel_importer import get_all_treatment_payments
    payments = get_all_treatment_payments()

    if search:
        payments = [p for p in payments if search in p.name]

    result = []
    for p in payments:
        result.append({
            'id': p.id,
            'name': p.name,
            'bank_account': p.bank_account,
            'amount': p.amount,
            'import_time': p.import_time,
        })

    return jsonify({'payments': result})


@app.route('/api/treatment_payment/<payment_id>', methods=['DELETE'])
@require_auth
def api_delete_treatment_payment(payment_id):
    if request.current_user.get('role') != 'admin':
        return jsonify({'error': '权限不足'}), 403

    from excel_importer import delete_treatment_payment
    try:
        delete_treatment_payment(payment_id)
        add_operation_log(request.current_user['id'], request.current_user['username'],
                         "删除待遇发放记录", f"ID:{payment_id}",
                         request.remote_addr, "Web")
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 400


@app.route('/api/treatment_payment/export')
@require_auth
def api_export_treatment_payment():
    search = request.args.get('search', None)

    from excel_importer import get_all_treatment_payments
    payments = get_all_treatment_payments()

    if search:
        payments = [p for p in payments if search in p.name]

    import pandas as pd
    import io
    from flask import Response
    data = []
    for p in payments:
        data.append({
            '姓名': p.name,
            '银行账号': p.bank_account,
            '金额': p.amount,
            '导入时间': p.import_time,
        })

    if not data:
        return jsonify({'error': '没有数据可导出'}), 400

    df = pd.DataFrame(data)
    output = io.BytesIO()
    df.to_excel(output, index=False, engine='openpyxl')
    output.seek(0)

    from datetime import datetime
    filename = f"待遇发放_{datetime.now().strftime('%Y%m%d')}.xlsx"

    add_operation_log(request.current_user['id'], request.current_user['username'],
                     "导出待遇发放", f"导出{len(data)}条记录",
                     request.remote_addr, "Web")

    return Response(
        output.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename="{filename}"'}
    )


def check_report_permission(report, user):
    if not report:
        return False, '报表不存在'
    
    username = user.get('username', '')
    is_admin = user.get('role') == 'admin'
    
    if is_admin:
        return True, ''
    
    if report.get('username') == username:
        return True, ''
    
    return False, '无权访问此报表'


@app.route('/api/reports/upload', methods=['POST'])
@require_auth
def api_report_upload():
    user = request.current_user
    username = user.get('username', '')
    if 'file' not in request.files:
        return jsonify({'error': '未选择文件'}), 400
    file = request.files['file']
    if not file.filename:
        return jsonify({'error': '未选择文件'}), 400

    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in ['.xlsx', '.xls', '.et']:
        return jsonify({'error': '仅支持Excel文件'}), 400

    report_name = request.form.get('name', file.filename)
    report_id = request.form.get('report_id', '')
    file_data = file.read()

    from report_storage import save_report, update_report, append_report_log, get_report_dir, find_report_by_id

    if report_id:
        existing = find_report_by_id(report_id)
        if existing and (existing['username'] == username or user.get('role') == 'admin'):
            meta = update_report(report_id, file_data, ext)
            if meta:
                add_report_log(report_id, user['id'], username, "更新报表",
                               f"文件名:{report_name}, 大小:{len(file_data)}字节",
                               request.remote_addr, "桌面端")
                log_entry = {
                    "time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    "user": username,
                    "action": "更新报表",
                    "detail": f"文件名:{report_name}, 大小:{len(file_data)}字节",
                    "ip": request.remote_addr
                }
                append_report_log(existing['server_path'], log_entry)
                return jsonify({'success': True, 'report_id': report_id, 'report_name': report_name})

    meta = save_report(username, report_name, file_data, ext)

    add_report_log(meta["report_id"], user['id'], username, "上传报表",
                   f"文件名:{report_name}, 大小:{len(file_data)}字节",
                   request.remote_addr, "桌面端")

    log_entry = {
        "time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "user": username,
        "action": "上传报表",
        "detail": f"文件名:{report_name}, 大小:{len(file_data)}字节",
        "ip": request.remote_addr
    }
    report_dir = get_report_dir(username, report_name)
    append_report_log(report_dir, log_entry)

    return jsonify({'success': True, 'report_id': meta['report_id'], 'report_name': report_name})


@app.route('/api/reports')
@require_auth
def api_report_list():
    from report_storage import list_all_reports
    reports = list_all_reports()

    result = []
    for r in reports:
        result.append({
            'report_id': r['report_id'],
            'username': r['username'],
            'original_name': r['original_name'],
            'upload_time': r['upload_time'],
            'file_size': r.get('file_size', 0),
        })
    return jsonify(result)


@app.route('/api/reports/<report_id>/download')
@require_auth
def api_report_download(report_id):
    user = request.current_user
    username = user.get('username', '')

    from report_storage import find_report_by_id, load_report, append_report_log
    report = find_report_by_id(report_id)
    
    has_perm, error_msg = check_report_permission(report, user)
    if not has_perm:
        return jsonify({'error': error_msg}), 403

    try:
        file_data = load_report(report['server_path'], report['stored_filename'])
    except Exception as e:
        return jsonify({'error': f'报表文件解密失败或已损坏：{str(e)}'}), 500

    add_report_log(report_id, user['id'], username, "下载报表",
                   f"文件名:{report['original_name']}",
                   request.remote_addr, "桌面端")

    log_entry = {
        "time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "user": username,
        "action": "下载报表",
        "detail": f"文件名:{report['original_name']}",
        "ip": request.remote_addr
    }
    append_report_log(report['server_path'], log_entry)

    from flask import Response
    return Response(
        file_data,
        mimetype='application/octet-stream',
        headers={'Content-Disposition': f'attachment; filename="{report["original_name"]}"'}
    )


@app.route('/api/reports/<report_id>/logs')
@require_auth
def api_report_logs(report_id):
    from report_storage import find_report_by_id
    report = find_report_by_id(report_id)
    
    has_perm, error_msg = check_report_permission(report, request.current_user)
    if not has_perm:
        return jsonify({'error': error_msg}), 403

    db_logs = get_report_logs(report_id)
    result = [{
        'time': l.created_at,
        'user': l.username,
        'action': l.operation_type,
        'detail': l.operation_detail,
        'ip': l.ip_address,
    } for l in db_logs]

    local_log_path = os.path.join(report['server_path'], "operation_log.json")
    local_logs = []
    if os.path.exists(local_log_path):
        with open(local_log_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
        local_logs = data.get('logs', [])

    return jsonify({'db_logs': result, 'file_logs': local_logs})


@app.route('/api/reports/<report_id>/edit')
@require_auth
def api_report_edit_page(report_id):
    from report_storage import find_report_by_id
    report = find_report_by_id(report_id)
    
    has_perm, error_msg = check_report_permission(report, request.current_user)
    if not has_perm:
        return jsonify({'error': error_msg}), 403
    
    return send_from_directory(STATIC_DIR, 'report_editor.html')


@app.route('/api/reports/<report_id>/data')
@require_auth
def api_report_data(report_id):
    from report_storage import find_report_by_id, load_report
    report = find_report_by_id(report_id)
    
    has_perm, error_msg = check_report_permission(report, request.current_user)
    if not has_perm:
        return jsonify({'error': error_msg}), 403

    try:
        file_data = load_report(report['server_path'], report['stored_filename'])
    except Exception as e:
        return jsonify({'error': f'报表文件解密失败或已损坏：{str(e)}'}), 500
    import io

    ext = report.get('file_ext', '')
    if ext == '.et':
        import xlrd
        wb = xlrd.open_workbook(file_contents=file_data)
        ws = wb.sheet_by_index(0)
        rows = []
        for i in range(ws.nrows):
            row = []
            for j in range(ws.ncols):
                cell = ws.cell(i, j)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        val = xlrd.xldate_as_datetime(cell.value, wb.datemode)
                        row.append(val.strftime("%Y-%m-%d"))
                    except Exception:
                        row.append(str(cell.value))
                elif cell.ctype == xlrd.XL_CELL_NUMBER:
                    if cell.value == int(cell.value):
                        row.append(str(int(cell.value)))
                    else:
                        row.append(str(cell.value))
                else:
                    row.append(str(cell.value) if cell.value else "")
            rows.append(row)
        return jsonify({
            'report_id': report_id,
            'report_name': report['original_name'],
            'username': report['username'],
            'headers': rows[0] if rows else [],
            'data': rows[1:] if len(rows) > 1 else [],
            'total_rows': len(rows),
        })

    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(file_data))
    ws = wb.active
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append([str(cell) if cell is not None else "" for cell in row])
    return jsonify({
        'report_id': report_id,
        'report_name': report['original_name'],
        'username': report['username'],
        'headers': rows[0] if rows else [],
        'data': rows[1:] if len(rows) > 1 else [],
        'total_rows': len(rows),
    })


@app.route('/api/reports/<report_id>/data', methods=['PUT'])
@require_auth
def api_report_save_data(report_id):
    user = request.current_user
    username = user.get('username', '')
    from report_storage import find_report_by_id, update_report, append_report_log, get_report_dir
    report = find_report_by_id(report_id)
    
    has_perm, error_msg = check_report_permission(report, user)
    if not has_perm:
        return jsonify({'error': error_msg}), 403

    try:
        data, err = _get_json()
        if err:
            return err
        headers = data.get('headers', [])
        rows = data.get('data', [])
        changes = data.get('changes', [])

        import io
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        if headers:
            ws.append(headers)
        for row in rows:
            ws.append(row)
        buf = io.BytesIO()
        wb.save(buf)
        file_data = buf.getvalue()

        ext = report.get('file_ext', '.xlsx')
        if ext and not ext.startswith('.'):
            ext = '.' + ext
        if ext == '.et':
            ext = '.xlsx'
        update_report(report_id, file_data, ext)

        if changes:
            for c in changes:
                add_report_log(report_id, user['id'], username, "修改单元格",
                               f"单元格:{c.get('cell','')}, 原值:{c.get('old','')}, 新值:{c.get('new','')}",
                               request.remote_addr, "Web")

            report_dir = report.get('server_path', '')
            if report_dir:
                log_entry = {
                    "time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    "user": username,
                    "action": "编辑保存",
                    "detail": f"修改了{len(changes)}个单元格",
                    "ip": request.remote_addr,
                }
                append_report_log(report_dir, log_entry)

        add_report_log(report_id, user['id'], username, "保存报表",
                       f"文件名:{report['original_name']}, {len(rows)}行数据",
                       request.remote_addr, "Web")

        return jsonify({'success': True})
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500


@app.route('/api/reports/<report_id>/import_ids', methods=['POST'])
@require_auth
def api_report_import_ids(report_id):
    user = request.current_user
    username = user.get('username', '')
    from report_storage import find_report_by_id, load_report, update_report, append_report_log, get_report_dir
    report = find_report_by_id(report_id)
    
    has_perm, error_msg = check_report_permission(report, user)
    if not has_perm:
        return jsonify({'error': error_msg}), 403

    data, err = _get_json()
    if err:
        return err
    id_list = data.get('ids', [])

    ext = report.get('file_ext', '.xlsx')
    file_data = load_report(report['server_path'], report['stored_filename'])

    import io
    if ext == '.et':
        import xlrd
        wb = xlrd.open_workbook(file_contents=file_data)
        ws = wb.sheet_by_index(0)
        rows = []
        for i in range(ws.nrows):
            row = []
            for j in range(ws.ncols):
                cell = ws.cell(i, j)
                if cell.ctype == xlrd.XL_CELL_NUMBER and cell.value == int(cell.value):
                    row.append(str(int(cell.value)))
                else:
                    row.append(str(cell.value) if cell.value else "")
            rows.append(row)
    else:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(file_data))
        ws = wb.active
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append([str(cell) if cell is not None else "" for cell in row])

    if not rows:
        return jsonify({'error': '报表无数据'}), 400

    headers = rows[0]
    id_card_col = -1
    phone_col = -1
    social1_col = -1
    social2_col = -1
    social3_col = -1
    for i, h in enumerate(headers):
        hl = h.lower().strip()
        if '身份证' in hl or 'id_card' in hl or '证件号' in hl:
            id_card_col = i
        elif '电话' in hl or 'phone' in hl or '手机' in hl:
            phone_col = i
        elif '社保卡号1' in hl or '社保卡1' in hl or (social1_col == -1 and ('社保' in hl or 'social' in hl)):
            social1_col = i
        elif '社保卡号2' in hl or '社保卡2' in hl:
            social2_col = i
        elif '社保卡号3' in hl or '社保卡3' in hl:
            social3_col = i

    if id_card_col == -1:
        for i, h in enumerate(headers):
            if i == phone_col or i == social1_col or i == social2_col or i == social3_col:
                continue
            id_card_col = i
            break

    if id_card_col == -1:
        return jsonify({'error': '报表中未找到身份证号列'}), 400

    id_map = {item['id_card']: item for item in id_list if item.get('id_card')}
    updated = 0
    not_found = 0
    for i in range(1, len(rows)):
        row = rows[i]
        cell_id = row[id_card_col].strip() if id_card_col < len(row) else ""
        if cell_id not in id_map:
            not_found += 1
            continue
        item = id_map[cell_id]
        if item.get('phone') and phone_col >= 0 and phone_col < len(row):
            row[phone_col] = item['phone']
        if item.get('social1') and social1_col >= 0 and social1_col < len(row):
            row[social1_col] = item['social1']
        if item.get('social2') and social2_col >= 0 and social2_col < len(row):
            row[social2_col] = item['social2']
        if item.get('social3') and social3_col >= 0 and social3_col < len(row):
            row[social3_col] = item['social3']
        updated += 1

    save_ext = ext if ext != '.et' else '.xlsx'
    import openpyxl
    save_wb = openpyxl.Workbook()
    save_ws = save_wb.active
    for row in rows:
        save_ws.append(row)
    buf = io.BytesIO()
    save_wb.save(buf)
    update_report(report_id, buf.getvalue(), save_ext)

    add_report_log(report_id, user['id'], username, "导入证件号码",
                   f"更新{updated}条,未找到{not_found}条",
                   request.remote_addr, "桌面端")
    log_entry = {
        "time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "user": username,
        "action": "导入证件号码",
        "detail": f"更新{updated}条,未找到{not_found}条",
        "ip": request.remote_addr
    }
    report_dir = report.get('server_path', '')
    if report_dir:
        append_report_log(report_dir, log_entry)

    return jsonify({'success': True, 'updated': updated, 'not_found': not_found})


@app.route('/api/backup', methods=['POST'])
@require_auth
def api_backup():
    if request.current_user.get('role') != 'admin':
        return jsonify({'error': '权限不足'}), 403
    from excel_importer import backup_database
    result = backup_database(manual=True)
    add_operation_log(request.current_user['id'], request.current_user['username'], "数据备份", f"成功:{result.get('success',False)}", request.remote_addr, "Web")
    return jsonify(result)


@app.route('/api/backup/download')
@require_auth
def api_backup_download():
    if request.current_user.get('role') != 'admin':
        return jsonify({'error': '权限不足'}), 403
    
    user_id = request.current_user['id']
    now = datetime.now().timestamp()
    if user_id in _backup_download_times:
        elapsed = now - _backup_download_times[user_id]
        if elapsed < 60:
            return jsonify({'error': f'下载过于频繁，请{int(60 - elapsed)}秒后重试'}), 429
    _backup_download_times[user_id] = now
    
    from config import BACKUP_DIR
    from excel_importer import get_backup_status
    status = get_backup_status()
    if not status.get('last_backup'):
        return jsonify({'error': '没有备份'}), 404
    backups = []
    for item in os.listdir(BACKUP_DIR):
        if item.startswith('backup_'):
            item_path = os.path.join(BACKUP_DIR, item)
            complete_file = os.path.join(item_path, 'backup_complete.txt')
            if os.path.exists(complete_file):
                backups.append({'path': item_path, 'time': os.path.getmtime(item_path)})
    if not backups:
        return jsonify({'error': '没有备份'}), 404
    latest = max(backups, key=lambda x: x['time'])
    db_path = os.path.join(latest['path'], 'population.db')
    if os.path.exists(db_path):
        safe_name = f'backup_{secrets.token_hex(8)}.db'
        add_operation_log(request.current_user['id'], request.current_user['username'], 
                         "下载备份", f"文件:{os.path.basename(latest['path'])}", 
                         request.remote_addr, "Web")
        return send_file(db_path, as_attachment=True, download_name=safe_name)
    return jsonify({'error': '备份文件不存在'}), 404


@app.route('/api/backups')
@require_auth
def api_backup_list():
    user = request.current_user
    if user.get('role') != 'admin':
        return jsonify({'error': '权限不足'}), 403
    
    from excel_importer import get_backup_list
    backups = get_backup_list()
    
    return jsonify({'backups': backups})


@app.route('/api/backup/restore', methods=['POST'])
@require_auth
def api_backup_restore():
    user = request.current_user
    if user.get('role') != 'admin':
        return jsonify({'error': '权限不足'}), 403
    
    data = request.get_json()
    backup_name = data.get('name', '')
    password = data.get('password', '')
    
    if not backup_name:
        return jsonify({'error': '缺少备份名称'}), 400
    
    if '..' in backup_name or '/' in backup_name or '\\' in backup_name:
        return jsonify({'error': '无效的备份名称'}), 400
    
    from config import BACKUP_DIR
    backup_path = os.path.join(BACKUP_DIR, backup_name)
    
    if not os.path.exists(backup_path):
        return jsonify({'error': '备份不存在'}), 404
    
    from excel_importer import restore_database
    result = restore_database(backup_path, password if password else None)
    
    if result.get('success'):
        add_operation_log(
            user['id'],
            user['username'],
            "恢复备份",
            result.get('timestamp', ''),
            request.remote_addr,
            "Web"
        )
        return jsonify(result)
    else:
        return jsonify(result), 400


def validate_date(date_str):
    if not date_str:
        return True
    try:
        datetime.strptime(date_str, '%Y-%m-%d')
        return True
    except ValueError:
        return False


@app.route('/api/logs')
@require_auth
def api_logs():
    user = request.current_user
    is_admin = user.get('role') == 'admin'
    username = user.get('username', '')
    
    limit = min(max(request.args.get('limit', 100, type=int), 1), 1000)
    
    filter_username = request.args.get('username', '')
    filter_operation_type = request.args.get('operation_type', '')
    filter_start_date = request.args.get('start_date', '')
    filter_end_date = request.args.get('end_date', '')
    
    if filter_start_date and not validate_date(filter_start_date):
        return jsonify({'error': '开始日期格式错误，应为 YYYY-MM-DD'}), 400
    
    if filter_end_date and not validate_date(filter_end_date):
        return jsonify({'error': '结束日期格式错误，应为 YYYY-MM-DD'}), 400
    
    if not is_admin:
        filter_username = username
    
    logs = get_operation_logs(
        limit=limit,
        username=filter_username if filter_username else None,
        operation_type=filter_operation_type if filter_operation_type else None,
        start_date=filter_start_date if filter_start_date else None,
        end_date=filter_end_date if filter_end_date else None
    )
    
    return jsonify([{
        'id': l.id,
        'username': l.username,
        'operation_type': l.operation_type,
        'operation_target': l.operation_target,
        'ip_address': l.ip_address,
        'client_type': l.client_type,
        'created_at': l.created_at,
        'is_high_risk': l.is_high_risk
    } for l in logs])


@app.route('/api/logs/types')
@require_auth
def api_log_types():
    types = get_operation_types()
    return jsonify({'types': types})


@app.route('/api/logs/stats')
@require_auth
def api_log_stats():
    user = request.current_user
    is_admin = user.get('role') == 'admin'
    
    if not is_admin:
        return jsonify({'error': '权限不足'}), 403
    
    stats = get_operation_log_stats()
    return jsonify(stats)


@app.route('/api/notifications', methods=['GET'])
@require_auth
def api_get_notifications():
    user = request.current_user
    username = user.get('username', '')
    limit = request.args.get('limit', 50, type=int)
    
    notifications = get_notifications(username=username, limit=limit)
    return jsonify({'notifications': notifications})


@app.route('/api/notifications/unread-count')
@require_auth
def api_get_unread_count():
    user = request.current_user
    username = user.get('username', '')
    
    count = get_unread_count(username)
    return jsonify({'count': count})


@app.route('/api/notifications/<notification_id>/read', methods=['POST'])
@require_auth
def api_mark_notification_read(notification_id):
    success = mark_notification_read(notification_id)
    if success:
        return jsonify({'success': True})
    else:
        return jsonify({'error': '通知不存在'}), 404


@app.route('/api/notifications/send', methods=['POST'])
@require_auth
def api_send_notification():
    user = request.current_user
    is_admin = user.get('role') == 'admin'
    
    if not is_admin:
        return jsonify({'error': '权限不足'}), 403
    
    data = request.get_json()
    if not data:
        return jsonify({'error': '无效请求数据'}), 400
    
    title = data.get('title', '').strip()
    content = data.get('content', '').strip()
    notification_type = data.get('notification_type', 'system')
    target_user = data.get('target_user', '').strip()
    
    if not title or not content:
        return jsonify({'error': '标题和内容不能为空'}), 400
    
    if notification_type not in ['system', 'business']:
        notification_type = 'system'
    
    sender_id = user.get('user_id', '')
    sender_name = user.get('real_name', user.get('username', ''))
    
    notification = create_notification(
        title=title,
        content=content,
        notification_type=notification_type,
        target_user=target_user,
        sender_id=sender_id,
        sender_name=sender_name
    )
    
    return jsonify({'success': True, 'notification': notification})


@app.route('/api/persons/batch-delete', methods=['POST'])
@require_auth
def api_batch_delete_persons():
    if request.current_user.get('role') != 'admin':
        return jsonify({'error': '权限不足'}), 403
    
    data = request.get_json()
    person_ids = data.get('ids', [])
    
    if not person_ids:
        return jsonify({'error': '未选择人员'}), 400
    
    session = get_session()
    try:
        # 支持人员ID和身份证号两种方式
        deleted_count = session.query(Person).filter(
            (Person.id.in_(person_ids)) | (Person.id_card.in_(person_ids))
        ).delete(synchronize_session='fetch')
        
        session.commit()
        
        add_operation_log(
            request.current_user['id'],
            request.current_user['username'],
            "批量删除人员",
            f"删除{deleted_count}人",
            request.remote_addr,
            "Web"
        )
        
        return jsonify({'success': True, 'deleted': deleted_count})
    except Exception as e:
        session.rollback()
        return jsonify({'error': str(e)}), 500
    finally:
        session.close()


@app.route('/api/persons/batch-update', methods=['POST'])
@require_auth
def api_batch_update_persons():
    """批量更新人员信息（仅管理员）"""
    if request.current_user.get('role') != 'admin':
        return jsonify({'error': '权限不足'}), 403
    
    data = request.get_json()
    person_ids = data.get('ids', [])
    updates = data.get('updates', {})
    
    if not person_ids:
        return jsonify({'error': '未选择人员'}), 400
    
    if not updates:
        return jsonify({'error': '未提供更新内容'}), 400
    
    session = get_session()
    try:
        persons = session.query(Person).filter(Person.id.in_(person_ids)).all()
        updated_count = 0
        
        for person in persons:
            for key, value in updates.items():
                if hasattr(person, key) and key not in ['id', 'id_card', 'hu_hao']:
                    setattr(person, key, value)
            person.updated_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            updated_count += 1
        
        session.commit()
        
        add_operation_log(
            request.current_user['id'],
            request.current_user['username'],
            "批量更新人员",
            f"更新{updated_count}人",
            request.remote_addr,
            "Web"
        )
        
        return jsonify({'success': True, 'updated': updated_count})
    except Exception as e:
        session.rollback()
        return jsonify({'error': str(e)}), 500
    finally:
        session.close()


@app.route('/api/students/batch-delete', methods=['POST'])
@require_auth
def api_batch_delete_students():
    """批量删除学生记录（仅管理员）"""
    if request.current_user.get('role') != 'admin':
        return jsonify({'error': '权限不足'}), 403
    
    data = request.get_json()
    student_ids = data.get('ids', [])
    
    if not student_ids:
        return jsonify({'error': '未选择学生'}), 400
    
    session = get_session()
    try:
        deleted_count = session.query(StudentEducation).filter(
            StudentEducation.id.in_(student_ids)
        ).delete(synchronize_session='fetch')
        
        session.commit()
        
        add_operation_log(
            request.current_user['id'],
            request.current_user['username'],
            "批量删除学生",
            f"删除{deleted_count}条记录",
            request.remote_addr,
            "Web"
        )
        
        return jsonify({'success': True, 'deleted': deleted_count})
    except Exception as e:
        session.rollback()
        return jsonify({'error': str(e)}), 500
    finally:
        session.close()
@app.route('/api/students/batch-update', methods=['POST'])
@require_auth
def api_batch_update_students():
    """批量更新学生信息（仅管理员）"""
    if request.current_user.get('role') != 'admin':
        return jsonify({'error': '权限不足'}), 403
    
    data = request.get_json()
    student_ids = data.get('ids', [])
    updates = data.get('updates', {})
    
    if not student_ids:
        return jsonify({'error': '未选择学生'}), 400
    
    if not updates:
        return jsonify({'error': '未提供更新内容'}), 400
    
    session = get_session()
    try:
        students = session.query(StudentEducation).filter(
            StudentEducation.id.in_(student_ids)
        ).all()
        updated_count = 0
        
        for student in students:
            for key, value in updates.items():
                if hasattr(student, key) and key != 'id':
                    setattr(student, key, value)
            updated_count += 1
        
        session.commit()
        
        add_operation_log(
            request.current_user['id'],
            request.current_user['username'],
            "批量更新学生",
            f"更新{updated_count}条记录",
            request.remote_addr,
            "Web"
        )
        
        return jsonify({'success': True, 'updated': updated_count})
    except Exception as e:
        session.rollback()
        return jsonify({'error': str(e)}), 500
    finally:
        session.close()


@app.route('/api/photos/batch-delete', methods=['POST'])
@require_auth
def api_batch_delete_photos():
    """批量删除照片（仅管理员）"""
    if request.current_user.get('role') != 'admin':
        return jsonify({'error': '权限不足'}), 403
    
    data = request.get_json()
    photos = data.get('photos', [])
    
    if not photos:
        return jsonify({'error': '未选择照片'}), 400
    
    deleted_count = 0
    not_found_count = 0
    errors = []
    
    session = get_session()
    try:
        for photo_info in photos:
            person_id = photo_info.get('person_id')
            photo_type = photo_info.get('photo_type', '身份证')
            id_card = photo_info.get('id_card')
            
            if person_id:
                # 使用 safe_id_card_query 支持身份证号查询
                p, error = safe_id_card_query(session, person_id)
                if error:
                    not_found_count += 1
                    continue
                id_card = p.id_card
                social1 = p.social_card1
                social2 = p.social_card2
                social3 = p.social_card3
            elif id_card:
                social1 = photo_info.get('social1', '')
                social2 = photo_info.get('social2', '')
                social3 = photo_info.get('social3', '')
            else:
                errors.append("缺少person_id或id_card")
                continue
            
            base_name = get_photo_base_name(photo_type, id_card, social1, social2, social3)
            if not base_name or base_name == "无":
                not_found_count += 1
                continue
            
            path = find_existing_photo(base_name, PHOTO_UPLOAD_DIR)
            if path and os.path.exists(path):
                try:
                    os.remove(path)
                    deleted_count += 1
                except Exception as e:
                    errors.append(f"删除失败: {path} - {str(e)}")
            else:
                not_found_count += 1
        
        add_operation_log(
            request.current_user['id'],
            request.current_user['username'],
            "批量删除照片",
            f"删除{deleted_count}张，未找到{not_found_count}张",
            request.remote_addr,
            "Web"
        )
        
        result = {
            'success': True,
            'deleted': deleted_count,
            'not_found': not_found_count
        }
        if errors:
            result['errors'] = errors
        
        return jsonify(result)
    except Exception as e:
        return jsonify({'error': str(e)}), 500
    finally:
        session.close()


@app.route('/api/photos/<id_card>')
@require_auth
def api_get_person_photos(id_card):
    """获取某人的所有照片"""
    session = get_session()
    try:
        p = session.query(Person).filter(Person.id_card == id_card).first()
        if not p:
            return jsonify({'error': '人员不存在'}), 404
        
        # 照片类型列表
        photo_types = ['身份证', '社保卡1', '社保卡2', '社保卡3']
        photos = []
        
        for photo_type in photo_types:
            base_name = get_photo_base_name(photo_type, p.id_card, p.social_card1, p.social_card2, p.social_card3)
            if base_name:
                path = find_existing_photo(base_name, PHOTO_UPLOAD_DIR)
                if path and os.path.exists(path):
                    photos.append({
                        'photo_type': photo_type,
                        'url': f'/api/photo/{os.path.basename(path)}',
                        'file_name': os.path.basename(path),
                        'exists': True
                    })
        
        return jsonify({'photos': photos})
    finally:
        session.close()


@app.route('/api/photo/<id_card>/<photo_type>', methods=['DELETE'])
@require_auth
def api_delete_photo_by_type(id_card, photo_type):
    """删除指定照片"""
    if request.current_user.get('role') not in ['admin', 'user']:
        return jsonify({'error': '权限不足'}), 403
    
    session = get_session()
    try:
        p = session.query(Person).filter(Person.id_card == id_card).first()
        if not p:
            return jsonify({'error': '人员不存在'}), 404
        
        # 生成照片文件名
        base_name = get_photo_base_name(photo_type, p.id_card, p.social_card1, p.social_card2, p.social_card3)
        if not base_name:
            return jsonify({'error': '无法定位照片'}), 400
        
        # 查找并删除照片
        path = find_existing_photo(base_name, PHOTO_UPLOAD_DIR)
        if not path or not os.path.exists(path):
            return jsonify({'error': '照片不存在'}), 404
        
        os.remove(path)
        
        add_operation_log(
            request.current_user['id'],
            request.current_user['username'],
            "删除照片",
            f"删除 {p.name} 的 {photo_type} 照片",
            request.remote_addr,
            "Web"
        )
        
        return jsonify({'success': True, 'message': '删除成功'})
    except Exception as e:
        return jsonify({'error': str(e)}), 500
    finally:
        session.close()


@app.route('/api/reports/custom', methods=['POST'])
@require_auth
def api_create_custom_report():
    """创建自定义报表"""
    user = request.current_user
    data, err = _get_json()
    if err:
        return err
    
    title = data.get('title', '').strip()
    report_type = data.get('report_type', '')
    dimensions = data.get('dimensions', [])
    date_range = data.get('date_range', {})
    chart_type = data.get('chart_type', 'bar')
    filters = data.get('filters', {})
    
    if not title:
        return jsonify({'error': '报表标题不能为空'}), 400
    
    if report_type not in ['person', 'student', 'subsidy']:
        return jsonify({'error': '无效的报表类型，支持: person, student, subsidy'}), 400
    
    if chart_type not in ['bar', 'line', 'pie', 'radar']:
        chart_type = 'bar'
    
    try:
        report = create_custom_report(
            user_id=user['id'],
            title=title,
            report_type=report_type,
            dimensions=dimensions,
            date_range=date_range,
            chart_type=chart_type,
            filters=filters
        )
        
        add_operation_log(
            user['id'],
            user['username'],
            "创建自定义报表",
            f"标题:{title}, 类型:{report_type}",
            request.remote_addr,
            "Web"
        )
        
        return jsonify({'success': True, 'report': report})
    except Exception as e:
        return jsonify({'error': str(e)}), 400


@app.route('/api/reports/custom')
@require_auth
def api_get_custom_reports():
    """获取报表列表"""
    user = request.current_user
    report_type = request.args.get('report_type', '')
    limit = request.args.get('limit', 50, type=int)
    
    # 普通用户只能看自己的报表，管理员可以看到所有
    user_id = None if user.get('role') == 'admin' else user['id']
    
    try:
        reports = get_custom_reports(user_id=user_id, report_type=report_type, limit=limit)
        return jsonify({'reports': reports})
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/api/reports/custom/<report_id>')
@require_auth
def api_get_custom_report(report_id):
    """获取报表详情"""
    user = request.current_user
    
    try:
        report = get_custom_report(report_id)
        if not report:
            return jsonify({'error': '报表不存在'}), 404
        
        # 权限检查：普通用户只能查看自己的报表
        if user.get('role') != 'admin' and report['user_id'] != user['id']:
            return jsonify({'error': '权限不足'}), 403
        
        return jsonify({'report': report})
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/api/reports/custom/<report_id>', methods=['PUT'])
@require_auth
def api_update_custom_report(report_id):
    """更新报表"""
    user = request.current_user
    data, err = _get_json()
    if err:
        return err
    
    try:
        # 先检查报表是否存在及权限
        existing = get_custom_report(report_id)
        if not existing:
            return jsonify({'error': '报表不存在'}), 404
        
        if user.get('role') != 'admin' and existing['user_id'] != user['id']:
            return jsonify({'error': '权限不足'}), 403
        
        # 验证报表类型
        if 'report_type' in data and data['report_type'] not in ['person', 'student', 'subsidy']:
            return jsonify({'error': '无效的报表类型'}), 400
        
        # 验证图表类型
        if 'chart_type' in data and data['chart_type'] not in ['bar', 'line', 'pie', 'radar']:
            data['chart_type'] = 'bar'
        
        updated = update_custom_report(report_id, **data)
        
        add_operation_log(
            user['id'],
            user['username'],
            "更新自定义报表",
            f"ID:{report_id}",
            request.remote_addr,
            "Web"
        )
        
        return jsonify({'success': True, 'report': updated})
    except Exception as e:
        return jsonify({'error': str(e)}), 400


@app.route('/api/reports/custom/<report_id>', methods=['DELETE'])
@require_auth
def api_delete_custom_report(report_id):
    """删除报表"""
    user = request.current_user
    
    try:
        # 先检查报表是否存在及权限
        existing = get_custom_report(report_id)
        if not existing:
            return jsonify({'error': '报表不存在'}), 404
        
        if user.get('role') != 'admin' and existing['user_id'] != user['id']:
            return jsonify({'error': '权限不足'}), 403
        
        deleted = delete_custom_report(report_id)
        
        if deleted:
            add_operation_log(
                user['id'],
                user['username'],
                "删除自定义报表",
                f"ID:{report_id}, 标题:{existing['title']}",
                request.remote_addr,
                "Web"
            )
            return jsonify({'success': True})
        else:
            return jsonify({'error': '删除失败'}), 500
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/api/reports/generate', methods=['POST'])
@require_auth
def api_generate_report():
    """生成报表数据（预览）"""
    user = request.current_user
    data, err = _get_json()
    if err:
        return err
    
    report_type = data.get('report_type', '')
    dimensions = data.get('dimensions', [])
    date_range = data.get('date_range', {})
    filters = data.get('filters', {})
    
    if report_type not in ['person', 'student', 'subsidy']:
        return jsonify({'error': '无效的报表类型，支持: person, student, subsidy'}), 400
    
    if not dimensions:
        return jsonify({'error': '请选择至少一个统计维度'}), 400
    
    try:
        result = generate_report_data(
            report_type=report_type,
            dimensions=dimensions,
            date_range=date_range,
            filters=filters
        )
        
        add_operation_log(
            user['id'],
            user['username'],
            "生成报表数据",
            f"类型:{report_type}, 维度:{','.join(dimensions)}",
            request.remote_addr,
            "Web"
        )
        
        return jsonify({'success': True, 'data': result})
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/api/check_token')
@require_auth
def api_check_token():
    """检查 Token 是否有效（移动端需要）"""
    return jsonify({
        'valid': True,
        'user': {
            'id': request.current_user['id'],
            'username': request.current_user['username'],
            'real_name': request.current_user.get('real_name', ''),
            'role': request.current_user['role']
        }
    })


@app.route('/api/change_password', methods=['POST'])
@require_auth
def api_change_password():
    """修改密码（移动端需要）"""
    ip = request.remote_addr
    is_allowed, remaining = check_sensitive_op_rate_limit(ip, 'change_password')
    if not is_allowed:
        return jsonify({'success': False, 'message': f'操作过于频繁，请{remaining}秒后重试'}), 429
    
    data, err = _get_json()
    if err:
        return err
    
    old_password = data.get('old_password', '')
    new_password = data.get('new_password', '')
    
    if not old_password or not new_password:
        return jsonify({'success': False, 'message': '请输入旧密码和新密码'}), 400
    
    valid, msg = validate_password(new_password)
    if not valid:
        return jsonify({'success': False, 'message': msg}), 400
    
    record_sensitive_op_attempt(ip, 'change_password')
    
    session = get_session()
    try:
        user = session.query(User).filter(User.id == request.current_user['id']).first()
        if not user:
            return jsonify({'success': False, 'message': '用户不存在'}), 404
        
        if not verify_password(old_password, user.password_hash):
            return jsonify({'success': False, 'message': '旧密码错误'}), 400
        
        user.password_hash = hash_password(new_password)
        session.commit()
        
        add_operation_log(
            user.id,
            user.username,
            "修改密码",
            "",
            request.remote_addr,
            "Mobile"
        )
        
        return jsonify({'success': True, 'message': '密码修改成功'})
    finally:
        session.close()


@app.route('/api/student/<id_card>')
@require_auth
def api_get_student(id_card):
    """获取单个学生详情"""
    session = get_session()
    try:
        student = session.query(StudentEducation).filter(
            StudentEducation.id_card == id_card
        ).first()
        
        if not student:
            return jsonify({'error': '学生不存在'}), 404
        
        person = session.query(Person).filter(Person.id == student.person_id).first()
        
        return jsonify({
            'student': {
                'id': student.id,
                'person_id': student.person_id,
                'id_card': student.id_card,
                'name': person.name if person else '',
                'school_name': student.school_name,
                'grade': student.grade,
                'grade_display': GRADE_DISPLAY.get(student.grade, student.grade),
                'class_name': student.class_name,
                'guardian': student.guardian,
                'guardian_phone': student.guardian_phone,
                'is_left_behind': student.is_left_behind,
                'student_status': student.student_status,
                'enrollment_date': student.enrollment_date,
                'graduation_date': student.graduation_date
            }
        })
    finally:
        session.close()


@app.route('/api/student', methods=['POST'])
@require_auth
def api_create_student():
    """添加学生"""
    if request.current_user.get('role') not in ['admin', 'user']:
        return jsonify({'error': '权限不足'}), 403
    
    data, err = _get_json()
    if err:
        return err
    
    id_card = data.get('id_card', '')
    if not id_card:
        return jsonify({'error': '身份证号不能为空'}), 400
    
    session = get_session()
    try:
        # 检查是否已存在
        existing = session.query(StudentEducation).filter(
            StudentEducation.id_card == id_card
        ).first()
        if existing:
            return jsonify({'error': '该身份证号的学生已存在'}), 400
        
        # 查找对应的人员记录
        person = session.query(Person).filter(Person.id_card == id_card).first()
        person_id = person.id if person else ''
        
        student = StudentEducation(
            id=f"edu_{secrets.token_hex(8)}",
            person_id=person_id,
            id_card=id_card,
            school_name=data.get('school_name', ''),
            grade=data.get('grade', ''),
            class_name=data.get('class_name', ''),
            guardian=data.get('guardian', ''),
            guardian_phone=data.get('guardian_phone', ''),
            is_left_behind=data.get('is_left_behind', '否'),
            student_status=data.get('student_status', '在读'),
            enrollment_date=data.get('enrollment_date', '')
        )
        
        session.add(student)
        session.commit()
        
        add_operation_log(
            request.current_user['id'],
            request.current_user['username'],
            "添加学生",
            id_card,
            request.remote_addr,
            "Mobile"
        )
        
        return jsonify({'success': True, 'id': student.id})
    except Exception as e:
        session.rollback()
        return jsonify({'error': str(e)}), 500
    finally:
        session.close()


@app.route('/api/student/<id_card>', methods=['PUT'])
@require_auth
def api_update_student(id_card):
    """更新学生信息"""
    if request.current_user.get('role') not in ['admin', 'user']:
        return jsonify({'error': '权限不足'}), 403
    
    data, err = _get_json()
    if err:
        return err
    
    session = get_session()
    try:
        student = session.query(StudentEducation).filter(
            StudentEducation.id_card == id_card
        ).first()
        
        if not student:
            return jsonify({'error': '学生不存在'}), 404
        
        # 更新允许修改的字段
        updatable_fields = [
            'school_name', 'grade', 'class_name', 
            'guardian', 'guardian_phone', 
            'is_left_behind', 'student_status', 'enrollment_date'
        ]
        
        for field in updatable_fields:
            if field in data:
                setattr(student, field, data[field])
        
        session.commit()
        
        add_operation_log(
            request.current_user['id'],
            request.current_user['username'],
            "更新学生信息",
            id_card,
            request.remote_addr,
            "Mobile"
        )
        
        return jsonify({'success': True})
    except Exception as e:
        session.rollback()
        return jsonify({'error': str(e)}), 500
    finally:
        session.close()


@app.route('/api/student/<id_card>', methods=['DELETE'])
@require_auth
def api_delete_student(id_card):
    """删除学生"""
    if request.current_user.get('role') not in ['admin', 'user']:
        return jsonify({'error': '权限不足'}), 403
    
    session = get_session()
    try:
        student = session.query(StudentEducation).filter(
            StudentEducation.id_card == id_card
        ).first()
        
        if not student:
            return jsonify({'error': '学生不存在'}), 404
        
        session.delete(student)
        session.commit()
        
        add_operation_log(
            request.current_user['id'],
            request.current_user['username'],
            "删除学生",
            id_card,
            request.remote_addr,
            "Mobile"
        )
        
        return jsonify({'success': True})
    except Exception as e:
        session.rollback()
        return jsonify({'error': str(e)}), 500
    finally:
        session.close()


if __name__ == '__main__':
    from database import is_database_encrypted, decrypt_database, encrypt_database, reset_engine, set_db_password, get_db_password, clear_db_password
    import atexit

    def _init_database():
        if is_database_encrypted():
            db_password = get_db_password()
            if db_password:
                if decrypt_database(db_password):
                    reset_engine()
                    return True
            if getattr(sys, 'frozen', False):
                from PyQt6.QtWidgets import QApplication, QInputDialog, QLineEdit
                _app = QApplication.instance() or QApplication(sys.argv)
                password, ok = QInputDialog.getText(
                    None, "数据库密码", "检测到加密数据库，请输入密码：",
                    QLineEdit.EchoMode.Password
                )
                if not ok or not password:
                    sys.exit(1)
                password = password.strip()
            else:
                print("检测到加密数据库，请输入数据库密码：")
                password = input("密码： ").strip()
            if not decrypt_database(password):
                if getattr(sys, 'frozen', False):
                    from PyQt6.QtWidgets import QMessageBox
                    QMessageBox.critical(None, "错误", "密码错误或数据库损坏！")
                else:
                    print("密码错误或数据库损坏！")
                sys.exit(1)
            reset_engine()
            set_db_password(password)
        init_db()

    if getattr(sys, 'frozen', False):
        from PyQt6.QtWidgets import QApplication, QFileDialog
        _app = QApplication.instance() or QApplication(sys.argv)
        if not os.path.exists(config.CONFIG_FILE):
            folder = QFileDialog.getExistingDirectory(
                None, "首次运行 - 请选择数据存储目录", config.BASE_DIR
            )
            if folder:
                os.makedirs(config.APPDATA_DIR, exist_ok=True)
                with open(config.CONFIG_FILE, "w", encoding="utf-8") as f:
                    f.write(config.encrypt_path(folder))
                config.DATA_DIR = folder
                config.DB_PATH = os.path.join(folder, "population.db")
                config.PHOTOS_DIR = os.path.join(folder, "Photos")
                config.BACKUP_DIR = os.path.join(folder, "backups")
                config.LOG_DIR = os.path.join(folder, "logs")
                config.CERT_DIR = os.path.join(folder, "certs")
            else:
                from PyQt6.QtWidgets import QMessageBox
                QMessageBox.critical(None, "错误", "未选择数据目录，程序退出")
                sys.exit(1)

    _init_database()

    def cleanup():
        pw = get_db_password()
        if pw:
            reset_engine()
            encrypt_database(pw)

    atexit.register(cleanup)

    from ssl_utils import generate_self_signed_cert
    cert_file, key_file = generate_self_signed_cert()
    use_https = bool(cert_file and key_file)

    if getattr(sys, 'frozen', False):
        import threading
        import socket
        import webbrowser
        from PyQt6.QtWidgets import (
            QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
            QLabel, QPushButton, QSystemTrayIcon, QMenu, QMessageBox
        )
        from PyQt6.QtGui import QIcon, QPixmap, QPainter, QColor, QFont
        from PyQt6.QtCore import Qt, QTimer, QThread, pyqtSignal

        def _get_local_ip():
            try:
                s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
                s.connect(("8.8.8.8", 80))
                ip = s.getsockname()[0]
                s.close()
                return ip
            except Exception:
                return "127.0.0.1"

        def _make_icon(color=QColor(0, 153, 76)):
            pixmap = QPixmap(64, 64)
            pixmap.fill(color)
            painter = QPainter(pixmap)
            painter.setPen(QColor(255, 255, 255))
            painter.setFont(QFont("Arial", 32, QFont.Weight.Bold))
            painter.drawText(pixmap.rect(), Qt.AlignmentFlag.AlignCenter, "S")
            painter.end()
            return QIcon(pixmap)

        def run_server():
            if use_https:
                app.run(host='::', port=5000, debug=False, ssl_context=(cert_file, key_file))
            else:
                app.run(host='::', port=5000, debug=False)

        tray_app = QApplication(sys.argv)
        tray_app.setQuitOnLastWindowClosed(False)

        local_ip = _get_local_ip()
        protocol = "https" if use_https else "http"
        lan_url = f"{protocol}://{local_ip}:5000"
        local_url = f"{protocol}://127.0.0.1:5000"

        window = QMainWindow()
        window.setWindowTitle("人口信息查询系统 - 服务器")
        window.setFixedSize(420, 280)
        window.setWindowIcon(_make_icon())

        central = QWidget()
        window.setCentralWidget(central)
        layout = QVBoxLayout(central)

        title = QLabel("人口信息查询系统 服务器端")
        title.setFont(QFont("微软雅黑", 14, QFont.Weight.Bold))
        title.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(title)

        layout.addSpacing(10)

        status_label = QLabel("状态：启动中...")
        status_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(status_label)

        url_label = QLabel(f"本机访问：{local_url}")
        url_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(url_label)

        lan_label = QLabel(f"局域网访问：{lan_url}")
        lan_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(lan_label)

        layout.addSpacing(15)

        btn_layout = QHBoxLayout()
        btn_open = QPushButton("打开浏览器")
        btn_open.setFixedHeight(36)
        btn_open.clicked.connect(lambda: webbrowser.open(local_url))
        btn_layout.addWidget(btn_open)

        btn_exit = QPushButton("停止服务")
        btn_exit.setFixedHeight(36)
        btn_exit.setStyleSheet("background-color: #C8373E; color: white;")
        btn_layout.addWidget(btn_exit)
        layout.addLayout(btn_layout)

        tray_icon = QSystemTrayIcon(_make_icon())
        tray_icon.setToolTip(f"人口信息查询系统服务器\n{lan_url}")

        tray_menu = QMenu()
        tray_menu.addAction("显示窗口", window.showNormal)
        tray_menu.addAction("打开浏览器", lambda: webbrowser.open(local_url))
        tray_menu.addSeparator()
        tray_menu.addAction("退出", tray_app.quit)
        tray_icon.setContextMenu(tray_menu)
        tray_icon.activated.connect(lambda reason: (window.showNormal(), window.activateWindow()) if reason == QSystemTrayIcon.ActivationReason.Trigger else None)
        tray_icon.show()

        class DownloadWorker(QThread):
            progress = pyqtSignal(int, int)
            finished = pyqtSignal(str, str)

            def __init__(self, version_info):
                super().__init__()
                self.version_info = version_info

            def run(self):
                def on_progress(downloaded, total):
                    self.progress.emit(downloaded, total)

                filepath, error = download_update(self.version_info, progress_callback=on_progress)
                self.finished.emit(filepath or "", error)

        def check_server_updates():
            """检查服务器更新"""
            from PyQt6.QtWidgets import QDialog, QVBoxLayout, QLabel, QPushButton, QHBoxLayout, QTextEdit, QProgressBar
            from PyQt6.QtGui import QFont

            dialog = QDialog()
            dialog.setWindowTitle("检查更新")
            dialog.setFixedSize(480, 320)
            dialog.setWindowFlags(dialog.windowFlags() & ~Qt.WindowType.WindowContextHelpButtonHint)
            layout = QVBoxLayout(dialog)
            layout.setSpacing(12)

            status_label = QLabel("正在检查更新...")
            status_label.setFont(QFont("Microsoft YaHei", 11))
            status_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
            layout.addWidget(status_label)

            progress = QProgressBar()
            progress.setRange(0, 0)
            layout.addWidget(progress)

            info_text = QTextEdit()
            info_text.setReadOnly(True)
            info_text.setMaximumHeight(120)
            info_text.setVisible(False)
            layout.addWidget(info_text)

            btn_layout = QHBoxLayout()
            btn_close = QPushButton("关闭")
            btn_close.clicked.connect(dialog.close)
            btn_close.setVisible(False)
            btn_layout.addStretch()
            btn_layout.addWidget(btn_close)
            layout.addLayout(btn_layout)

            dialog.show()
            QApplication.processEvents()

            version_info, error = check_update("server")

            progress.setRange(0, 1)
            progress.setValue(1)

            if error:
                status_label.setText("检查更新失败")
                status_label.setStyleSheet("color: red;")
                info_text.setVisible(True)
                info_text.setText(error)
                btn_close.setVisible(True)
                return

            if not version_info:
                status_label.setText("当前已是最新版本")
                status_label.setStyleSheet("color: green;")
                btn_close.setVisible(True)
                return

            latest_version = version_info.get("server_version", version_info.get("latest_version", ""))
            release_notes = version_info.get("release_notes", "无")
            file_size = version_info.get("server_file_size", version_info.get("file_size", 0))

            status_label.setText(f"发现新版本：{latest_version}")
            status_label.setStyleSheet("color: #0066CC; font-weight: bold;")

            size_text = ""
            if file_size > 1024 * 1024:
                size_text = f"文件大小：{file_size / 1024 / 1024:.1f} MB"
            elif file_size > 1024:
                size_text = f"文件大小：{file_size / 1024:.1f} KB"

            info_text.setVisible(True)
            info_text.setText(f"更新内容：\n{release_notes}\n\n{size_text}")
            info_text.setStyleSheet("background-color: #F5F5F5;")

            btn_close.setText("暂不更新")
            btn_close.setVisible(True)

            btn_update = QPushButton("立即更新")
            btn_update.setStyleSheet("background-color: #0078D4; color: white; font-weight: bold; padding: 8px 24px;")
            btn_layout.insertWidget(0, btn_update)

            def do_update():
                btn_update.setEnabled(False)
                btn_update.setText("下载中...")
                btn_close.setEnabled(False)
                progress.setRange(0, 100)
                progress.setValue(0)
                status_label.setText("正在下载更新...")
                status_label.setStyleSheet("color: #0066CC;")
                info_text.setVisible(False)

                dialog._worker = DownloadWorker(version_info)

                def on_progress(downloaded, total):
                    if total > 0:
                        percent = int(downloaded * 100 / total)
                        progress.setValue(percent)
                        mb_done = downloaded / 1024 / 1024
                        mb_total = total / 1024 / 1024
                        status_label.setText(f"正在下载更新... {mb_done:.1f}/{mb_total:.1f} MB ({percent}%)")
                    else:
                        mb_done = downloaded / 1024 / 1024
                        status_label.setText(f"正在下载更新... {mb_done:.1f} MB")

                def on_finished(filepath, error):
                    if error:
                        progress.setRange(0, 1)
                        progress.setValue(1)
                        status_label.setText("下载失败")
                        status_label.setStyleSheet("color: red;")
                        info_text.setVisible(True)
                        info_text.setText(error)
                        btn_update.setText("下载失败")
                        btn_close.setEnabled(True)
                        return

                    status_label.setText("正在验证文件...")
                    progress.setRange(0, 0)

                    checksum = version_info.get("server_checksum", version_info.get("checksum", ""))
                    if not verify_file(filepath, checksum):
                        progress.setRange(0, 1)
                        progress.setValue(1)
                        status_label.setText("验证失败")
                        status_label.setStyleSheet("color: red;")
                        info_text.setVisible(True)
                        info_text.setText("下载文件验证失败，请重新下载。")
                        os.remove(filepath)
                        btn_update.setText("验证失败")
                        btn_close.setEnabled(True)
                        return

                    status_label.setText("正在准备安装...")
                    update_local_version(version_info)

                    ok, err = apply_update(filepath)
                    if ok:
                        status_label.setText("更新成功，程序即将重启...")
                        status_label.setStyleSheet("color: green; font-weight: bold;")
                        progress.setRange(0, 1)
                        progress.setValue(1)
                        btn_update.setVisible(False)
                        btn_close.setText("立即重启")
                        btn_close.setVisible(True)
                        btn_close.setEnabled(True)
                        btn_close.clicked.disconnect()
                        btn_close.clicked.connect(tray_app.quit)
                    else:
                        progress.setRange(0, 1)
                        progress.setValue(1)
                        status_label.setText("下载完成")
                        status_label.setStyleSheet("color: green;")
                        info_text.setVisible(True)
                        info_text.setText(f"更新文件已下载到：\n{filepath}\n\n请手动安装更新。")
                        btn_update.setVisible(False)
                        btn_close.setEnabled(True)

                dialog._worker.progress.connect(on_progress)
                dialog._worker.finished.connect(on_finished)
                dialog._worker.start()

            btn_update.clicked.connect(do_update)

        def start_server():
            server_thread = threading.Thread(target=run_server, daemon=True)
            server_thread.start()
            status_label.setText("状态：运行中")
            status_label.setStyleSheet("color: green; font-weight: bold;")
            # 启动后检查更新
            QTimer.singleShot(2000, check_server_updates)

        btn_exit.clicked.connect(tray_app.quit)
        window.closeEvent = lambda e: (e.ignore(), window.hide())

        QTimer.singleShot(100, start_server)
        window.show()

        sys.exit(tray_app.exec())
    else:
        print("=" * 60)
        print("人口信息查询系统 - 安全启动")
        print("=" * 60)
        
        # 强制HTTPS检查
        force_https = os.environ.get('FORCE_HTTPS', '1') == '1'
        
        if not use_https and force_https:
            print("\n" + "!" * 60)
            print("安全警告: HTTPS不可用!")
            print("!" * 60)
            print("\n正在尝试生成自签名证书...")
            try:
                os.makedirs(config.CERT_DIR, exist_ok=True)
                from ssl_utils import generate_self_signed_cert
                cert_file, key_file = generate_self_signed_cert()
                if cert_file and key_file:
                    use_https = True
                    print(f"✓ 自签名证书生成成功")
                    print(f"  证书位置: {cert_file}")
                else:
                    print("✗ 自签名证书生成失败")
                    print("\n请手动生成SSL证书:")
                    print("  openssl req -x509 -newkey rsa:2048 -keyout server.key -out server.crt -days 365 -nodes")
                    print("\n或设置环境变量 FORCE_HTTPS=0 允许HTTP运行（不推荐）")
                    sys.exit(1)
            except Exception as e:
                print(f"✗ 证书生成失败: {e}")
                sys.exit(1)
        
        if use_https:
            print(f"\n✓ HTTPS已启用")
            print(f"  访问地址: https://127.0.0.1:5000")
            print(f"  证书文件: {cert_file}")
            print("\n⚠ 注意: 使用自签名证书，浏览器会提示不安全")
            print("  这是正常的，点击'继续访问'即可")
            print("\n⚠ 生产环境建议:")
            print("  1. 使用正式SSL证书（Let's Encrypt免费）")
            print("  2. 配置nginx反向代理")
        else:
            print("\n" + "!" * 60)
            print("警告: 使用HTTP运行，所有数据明文传输！")
            print("!" * 60)
            print(f"  访问地址: http://127.0.0.1:5000")
            print("\n强烈建议启用HTTPS以保护数据安全")
        
        print("=" * 60)
        
        if use_https:
            app.run(host='::', port=5000, debug=False, ssl_context=(cert_file, key_file))
        else:
            app.run(host='::', port=5000, debug=False)
