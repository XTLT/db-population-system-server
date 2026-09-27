def export_student_education_template(file_path: str):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "学生就读信息"

    rules = [
        ("身份证号", "18位身份证号码，必须在系统中存在"),
        ("学校名称", "学校全称"),
        ("年级", "可选值：T1（托班）、K1-K3（幼儿园）、G1-G12（小学至高中）、V1-V3（职高）、J1-J3（技校）、C1-C3（中专）、Z1-Z3（专科）、U1-U5（大学）、M1-M3（硕士）、D1-D4（博士）、GRADUATED（已毕业）"),
        ("班级", "班级名称，如：1班、2班"),
        ("监护人", "监护人姓名，可选择本户成员或手动输入"),
        ("监护人电话", "监护人联系电话"),
        ("是否留守儿童", "可选值：是、否"),
        ("学生状态", "可选值：在读、留级、跳级、休学、转学、毕业"),
        ("入学日期", "格式：yyyy-MM-dd，如：2024-09-01")
    ]

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

    for col_idx, (field, rule) in enumerate(rules, 1):
        cell = ws.cell(row=1, column=col_idx, value=field)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

        rule_cell = ws.cell(row=2, column=col_idx, value=rule)
        rule_cell.font = Font(color="666666")
        rule_cell.alignment = Alignment(wrap_text=True)

    ws.row_dimensions[1].height = 25
    ws.row_dimensions[2].height = 40

    for col_idx in range(1, len(rules) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 20

    wb.save(file_path)
