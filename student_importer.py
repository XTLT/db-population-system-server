import os
import pandas as pd
from datetime import datetime
import shutil

from config import DB_PATH, BACKUP_DIR, STUDENT_EDUCATION_COLUMNS, GRADES
from database import get_session, StudentEducation, EducationHistory, Person
from utils import safe_str


def _read_excel_robust(file_path: str, header=None):
    """读取 Excel 文件，自动检测格式和引擎"""
    ext = os.path.splitext(file_path)[1].lower()
    engines = (['openpyxl', 'xlrd'] if ext == '.xlsx'
               else ['xlrd', 'openpyxl'] if ext == '.xls'
               else ['openpyxl', 'xlrd'])
    for engine in engines:
        try:
            return pd.read_excel(file_path, dtype=str, header=header, engine=engine)
        except Exception:
            continue
    try:
        return pd.read_csv(file_path, dtype=str, header=header, encoding='utf-8')
    except Exception:
        return pd.read_csv(file_path, dtype=str, header=header, encoding='gbk')


def backup_database():
    if not os.path.exists(DB_PATH):
        return None
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = os.path.join(BACKUP_DIR, f"population_{timestamp}.db")
    shutil.copy2(DB_PATH, backup_path)
    return backup_path


def import_student_education(file_path: str, mode: str = "incremental", progress_callback=None) -> dict:
    backup_database()
    df = _read_excel_robust(file_path)
    df = df.fillna("")
    total = len(df)
    imported = 0
    updated = 0
    skipped = 0
    errors = []

    session = get_session()
    try:
        for idx, row in df.iterrows():
            try:
                id_card = safe_str(row.get("身份证号", ""))
                if not id_card:
                    skipped += 1
                    errors.append(f"行{idx + 2}: 身份证号为空")
                    continue

                person = session.query(Person).filter(Person.id_card == id_card).first()
                if not person:
                    skipped += 1
                    errors.append(f"行{idx + 2}: 身份证号 {id_card} 不存在")
                    continue

                grade = safe_str(row.get("年级", ""))
                if grade and grade not in GRADES:
                    skipped += 1
                    errors.append(f"行{idx + 2}: 年级 {grade} 不合法")
                    continue

                existing = session.query(StudentEducation).filter(
                    StudentEducation.person_id == person.id
                ).first()

                if existing:
                    old_grade = existing.grade
                    existing.school_name = safe_str(row.get("学校名称", existing.school_name))
                    existing.grade = grade or existing.grade
                    existing.class_name = safe_str(row.get("班级", existing.class_name))
                    existing.guardian = safe_str(row.get("监护人", existing.guardian))
                    existing.guardian_phone = safe_str(row.get("监护人电话", existing.guardian_phone))
                    existing.is_left_behind = safe_str(row.get("是否留守儿童", existing.is_left_behind))
                    existing.student_status = safe_str(row.get("学生状态", existing.student_status))
                    existing.enrollment_date = safe_str(row.get("入学日期", existing.enrollment_date))
                    updated += 1

                    if old_grade != existing.grade:
                        history = EducationHistory(
                            id=f"hist_{datetime.now().strftime('%Y%m%d%H%M%S')}_{person.id}",
                            person_id=person.id,
                            id_card=id_card,
                            school_name=existing.school_name,
                            grade=old_grade,
                            class_name=existing.class_name,
                            student_status=existing.student_status,
                            change_date=datetime.now().strftime("%Y-%m-%d"),
                            change_reason="Excel导入更新"
                        )
                        session.add(history)
                else:
                    edu = StudentEducation(
                        id=f"edu_{person.id}",
                        person_id=person.id,
                        id_card=id_card,
                        school_name=safe_str(row.get("学校名称", "")),
                        grade=grade,
                        class_name=safe_str(row.get("班级", "")),
                        guardian=safe_str(row.get("监护人", "")),
                        guardian_phone=safe_str(row.get("监护人电话", "")),
                        is_left_behind=safe_str(row.get("是否留守儿童", "否")),
                        student_status=safe_str(row.get("学生状态", "在读")),
                        enrollment_date=safe_str(row.get("入学日期", ""))
                    )
                    session.add(edu)
                    imported += 1

                if progress_callback:
                    progress_callback(int((idx + 1) / total * 100))
            except Exception as e:
                skipped += 1
                errors.append(f"行{idx + 2}: {str(e)}")

        session.commit()
    except Exception:
        session.rollback()
        raise
    finally:
        session.close()

    return {
        "imported": imported,
        "updated": updated,
        "skipped": skipped,
        "errors": errors,
        "total": total
    }


def export_student_education_template(file_path: str):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "学生就读信息"

    rules = [
        ("身份证号", "18位身份证号码，必须在系统中存在"),
        ("学校名称", "学校全称"),
        ("年级", "可选值：T1（托班）、K1-K3（幼儿园）、G1-G12（小学至高中）、V1-V3（职高）、J1-J3（技校）、C1-C3（中专）、Z1-Z3（专科）、U1-U5（大学）、M1-M3（硕士）、D1-D4（博士）、GRADUATED（已毕业）"),
        ("班级", "班级名称，如：1班、2班"),
        ("监护人", "监护人姓名，可选择本户成员或手动输入"),
        ("监护人电话", "监护人联系电话"),
        ("是否留守儿童", "可选值：是、否"),
        ("学生状态", "可选值：在读、留级、跳级、休学、转学、毕业"),
        ("入学日期", "格式：yyyy-MM-dd，如：2024-09-01")
    ]

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

    for col_idx, (field, rule) in enumerate(rules, 1):
        cell = ws.cell(row=1, column=col_idx, value=field)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

        rule_cell = ws.cell(row=2, column=col_idx, value=rule)
        rule_cell.font = Font(color="666666")
        rule_cell.alignment = Alignment(wrap_text=True)

    ws.row_dimensions[1].height = 25
    ws.row_dimensions[2].height = 40

    for col_idx in range(1, len(rules) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 20

    wb.save(file_path)


def export_student_education(file_path: str, progress_callback=None) -> bool:
    session = get_session()
    try:
        students = session.query(StudentEducation).all()
        if not students:
            return False

        data = []
        for i, s in enumerate(students):
            person = session.query(Person).filter(Person.id == s.person_id).first()
            data.append({
                "身份证号": s.id_card,
                "姓名": person.name if person else "",
                "学校名称": s.school_name,
                "年级": s.grade,
                "班级": s.class_name,
                "监护人": s.guardian,
                "监护人电话": s.guardian_phone,
                "是否留守儿童": s.is_left_behind,
                "学生状态": s.student_status,
                "入学日期": s.enrollment_date
            })
            if progress_callback and (i + 1) % 50 == 0:
                progress_callback(int((i + 1) / len(students) * 100))

        df = pd.DataFrame(data, columns=["身份证号", "姓名"] + STUDENT_EDUCATION_COLUMNS[1:])
        df.to_excel(file_path, index=False, engine="openpyxl")
        return True
    finally:
        session.close()
